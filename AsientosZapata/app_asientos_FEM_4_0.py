import streamlit as st
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

from funciones_asientos import (
    OPENSEES_DISPONIBLE, MEFError, NU_MAX, MAX_NODOS_MEF, TIPOS_SUELO,
    TENSION_CONSISTENTE, TENSION_HOLL, CRITERIOS_GOBERNANTE,
    holl_centro, calcular_steinbrenner, calcular_ec68, calcular_opensees_3d,
    sigma_v0, z_influencia_ec7, validar_estratigrafia, estimar_malla_3d,
    asiento_gobernante, estratos_cohesivos,
)


# ══════════════════════════════════════════════════════════════════════════
# INFORME WORD ESTÉTICO
# ══════════════════════════════════════════════════════════════════════════
def _fig_bytes(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=250, bbox_inches='tight')
    buf.seek(0)
    return buf

def _add_styled_table(doc, df, title):
    if title:
        h = doc.add_heading(title, level=2)
        if h.runs:
            h.runs[0].font.color.rgb = RGBColor(31, 73, 125)
    
    df = df.astype(str)
    table = doc.add_table(rows=1+len(df), cols=len(df.columns))
    table.style = 'Light Shading Accent 1' 
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = table.rows[0].cells
    for i, column in enumerate(df.columns):
        hdr_cells[i].text = column
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(8) 
                run.font.color.rgb = RGBColor(23, 54, 93) 
    
    for i, row in enumerate(df.itertuples(index=False)):
        row_cells = table.rows[i+1].cells
        for j, value in enumerate(row):
            row_cells[j].text = str(value)
            row_cells[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in row_cells[j].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()

def _body(doc, text, italic=False, gray=False, size=10, bold=False):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = par.add_run(text)
    r.font.size = Pt(size); r.font.italic = italic; r.font.bold = bold
    if gray:
        r.font.color.rgb = RGBColor(89, 89, 89)
    return par

def _bullets(doc, items):
    for it in items:
        par = doc.add_paragraph(style='List Bullet')
        r = par.add_run(it); r.font.size = Pt(10)

def _h2(doc, text):
    h = doc.add_heading(text, level=2)
    if h.runs:
        h.runs[0].font.color.rgb = RGBColor(31, 73, 125)
    return h

def _nota(doc, text):
    """Párrafo de aviso resaltado (fondo suave mediante borde-color de texto)."""
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = par.add_run(text)
    r.font.size = Pt(10); r.font.bold = True
    r.font.color.rgb = RGBColor(192, 0, 0)
    return par


def _fmt_mm(v, dec=3):
    return f"{v*1000:.{dec}f} mm" if (v is not None and np.isfinite(v)) else "—"


def tabla_comparativa(c):
    return pd.DataFrame({
        "Capa": c["df_st"]["Capa"],
        "z Techo [m]": c["df_st"]["z Techo [m]"],
        "z Base [m]": c["df_st"]["z Base [m]"],
        "Δs Steinbrenner [mm]": c["df_st"]["Δs [mm]"],
        "Δs Ec. Elástica [mm]": c["df_ec"]["Δs [mm]"].values,
        "Δs OpenSees 3D [mm]": c["df_mef"]["Δs [mm]"].values,
    })


def figura_bulbo(c, z_top, figsize=(5, 7)):
    """Tensiones bajo el centro frente a 0,20·σ′v0."""
    z_vals = np.linspace(0.05, z_top, 200)
    sz_v, sx_v, sy_v = holl_centro(c["p"], c["B"], c["L"], z_vals)
    sv_v = sigma_v0(z_vals, c["df"], c["NF"], c["D"], c["gamma_D"], c["gamma_sat_D"])
    fig, ax = plt.subplots(figsize=figsize)
    ax.plot(sz_v, z_vals, label=r"Vertical $\Delta\sigma_z$", color='red', lw=2)
    ax.plot(sx_v, z_vals, label=r"Horiz. trans. $\Delta\sigma_x$ (Holl, $\nu$=0,5)", color='blue', ls='--', lw=1)
    ax.plot(sy_v, z_vals, label=r"Horiz. long. $\Delta\sigma_y$ (Holl, $\nu$=0,5)", color='purple', ls='-.', lw=1)
    ax.plot(sv_v, z_vals, label=r"$\sigma'_{v0}$", color='saddlebrown', ls=':', lw=1.5)
    ax.plot(0.20*sv_v, z_vals, label=r"$0.20\,\sigma'_{v0}$ (criterio EC7)", color='green', lw=2)
    if c["zi"] <= z_top:
        ax.axhline(y=c["zi"], color='orange', ls='--', lw=1.5, label=f'z_i EC7 = {c["zi"]:.2f} m')
    nf_loc = c["NF"] - c["D"]
    if 0 < nf_loc < z_top:
        ax.axhline(y=nf_loc, color='deepskyblue', ls='-.', lw=1.2, label=f'NF ({nf_loc:.2f} m bajo la base)')
    ax.set_ylim(z_top, 0); ax.set_xlim(left=0)
    ax.set_xlabel("Tensión (kPa)"); ax.set_ylabel("Profundidad bajo la base z (m)")
    ax.set_title("Bulbo de presiones — centro de la cimentación")
    ax.legend(fontsize=8, loc='lower right')
    ax.grid(True, linestyle=':', alpha=0.4)
    ax.spines[['top', 'right']].set_visible(False)
    plt.tight_layout()
    return fig


def generar_word(c, fig_bulbo_bytes, meta=None):
    """Memoria de cálculo a partir de la instantánea `c` del cálculo realizado."""
    fecha = datetime.now().strftime("%d/%m/%Y — %H:%M")
    meta = meta or {}
    obra         = meta.get("obra") or "[Denominación de la obra]"
    peticionario = meta.get("peticionario") or "[Peticionario]"
    referencia   = meta.get("referencia") or "GEO-XXXX-ASN-01"
    autor        = meta.get("autor") or "Dpto. Geotecnia"
    revision     = meta.get("revision") or "00 — Emisión inicial"

    B, L, p, NF, D = c["B"], c["L"], c["p"], c["NF"], c["D"]
    z_max, zi = c["z_max"], c["zi"]
    s_st, s_ec, s_mef = c["tot_st"]*1000, c["tot_ec"]*1000, c["tot_mef"]*1000
    mef_ok = np.isfinite(s_mef)
    vals = [v for v in (s_st, s_ec, s_mef) if np.isfinite(v)]
    media = sum(vals) / len(vals)
    dispersion = (max(vals) - min(vals)) / max(media, 1e-9) * 100.0
    s_gob, origen = asiento_gobernante(c["tot_st"], c["tot_ec"], c["tot_mef"],
                                       c["criterio"], c["factor_rigidez"])
    LIM_ELS = float(c["s_adm"])
    cumple_els = s_gob <= LIM_ELS
    consistente = c["modo_tension"] == TENSION_CONSISTENTE
    info = c["mef_info"]

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
        fp = sec.footer.paragraphs[0]
        fp.text = f"Memoria de Cálculo de Asientos · {referencia} · Generado el {fecha}"
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.runs[0].font.size = Pt(8); fp.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(10)
    h1f = doc.styles['Heading 1'].font
    h1f.name = 'Calibri Light'; h1f.size = Pt(14); h1f.color.rgb = RGBColor(23, 54, 93); h1f.bold = True

    # ───────────────────────── PORTADA ─────────────────────────
    doc.add_paragraph(); doc.add_paragraph()
    pt = doc.add_paragraph(); pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt = pt.add_run('MEMORIA DE CÁLCULO DE ASIENTOS')
    rt.bold = True; rt.font.size = Pt(24); rt.font.color.rgb = RGBColor(23, 54, 93)
    ps = doc.add_paragraph(); ps.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = ps.add_run('Cimentación superficial · Métodos analíticos y MEF 3D')
    rs.font.size = Pt(13); rs.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph(); doc.add_paragraph()

    tid = doc.add_table(rows=6, cols=2); tid.style = 'Light Shading Accent 1'
    id_data = [('Obra / Proyecto', obra), ('Peticionario', peticionario),
               ('Referencia', referencia), ('Redactado por', autor),
               ('Fecha', fecha), ('Revisión', revision)]
    for i, (k, v) in enumerate(id_data):
        cs = tid.rows[i].cells
        cs[0].text = k; cs[1].text = str(v)
        cs[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cs[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if cs[0].paragraphs[0].runs:
            cs[0].paragraphs[0].runs[0].font.bold = True
    doc.add_page_break()

    # ───────────────── 1. OBJETO Y ALCANCE ─────────────────
    doc.add_heading('1. Objeto y alcance', level=1)
    _h2(doc, '1.1 Objeto')
    _body(doc, "El presente documento recoge el cálculo del asiento en el centro de una cimentación superficial "
               "rectangular sometida a carga vertical uniforme, mediante tres formulaciones: el método analítico "
               "de Steinbrenner, la integración numérica de la deformación vertical elástica y un modelo "
               "tridimensional de elementos finitos (OpenSeesPy).")
    _h2(doc, '1.2 Alcance: asiento elástico / inmediato')
    _body(doc, "Los tres métodos parten de la teoría de la elasticidad lineal. En consecuencia, el resultado "
               "corresponde al ASIENTO ELÁSTICO, obtenido con los módulos introducidos. Si estos son no drenados, "
               "el asiento es inmediato y no incluye la consolidación; si son drenados (E′), representa el asiento "
               "diferido en condiciones drenadas. El valor calculado corresponde al centro de un área flexible.")
    doc.add_page_break()

    # ───────────────── 2. NORMATIVA ─────────────────
    doc.add_heading('2. Normativa y referencias', level=1)
    _bullets(doc, [
        "CTE DB-SE-C «Seguridad estructural – Cimientos». Estados Límite de Servicio (asientos admisibles).",
        "UNE-EN 1997-1 (Eurocódigo 7), apartado 6.6.2: criterio de profundidad de influencia (Δσz ≤ 0,20·σ′v0).",
        "Steinbrenner (1934) / Bowles: desplazamiento vertical del semiespacio elástico bajo la esquina de un "
        "rectángulo; asiento de cada estrato como diferencia entre techo y base (aproximación de estrato finito).",
        "Holl (1940): tensiones bajo carga rectangular uniforme; las expresiones de σx y σy corresponden a ν = 0,5.",
        "Boussinesq: suma de tensiones normales θ = (1+ν)·P·z/(π·R³), válida para ν arbitrario.",
        "OpenSeesPy: biblioteca de elementos finitos; hexaedro de 8 nodos con formulación B-bar (bbarBrick).",
    ])
    doc.add_page_break()

    # ───────────────── 3. DATOS DE PARTIDA ─────────────────
    doc.add_heading('3. Datos de partida', level=1)
    _h2(doc, '3.1 Geometría, acciones y profundidades')
    d3 = [('Dimensiones en planta (B × L)', f'{B:.2f} m × {L:.2f} m'),
          ('Presión neta de trabajo (p)', f'{p:.1f} kPa'),
          ('Profundidad de cimentación (D)', f'{D:.2f} m'),
          ('Terreno sobre la base (γ / γsat)', f'{c["gamma_D"]:.1f} / {c["gamma_sat_D"]:.1f} kN/m³'),
          ('Nivel freático (NF, desde superficie)', f'{NF:.1f} m' if NF < 100 else 'No detectado'),
          ('Profundidad de influencia (z_i, EC7)', f'{zi:.2f} m bajo la base'
                                                   + ('' if c["zi_ok"] else ' (criterio no alcanzado)')),
          ('Profundidad de corte evaluada (z_max)', f'{z_max:.2f} m bajo la base'),
          ('Punto de cálculo', 'Centro de la zapata (superposición ×4 de cuadrantes)')]
    t3 = doc.add_table(rows=len(d3), cols=2); t3.style = 'Light Shading Accent 1'
    for i, (k, v) in enumerate(d3):
        cs = t3.rows[i].cells; cs[0].text = k; cs[1].text = v
        cs[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER; cs[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if cs[0].paragraphs[0].runs: cs[0].paragraphs[0].runs[0].font.bold = True
    _body(doc, "La presión p es la presión NETA de trabajo (descontada la sobrecarga de tierras retirada). "
               "Las profundidades z se miden desde la base de la zapata; D y NF, desde la superficie del terreno. "
               "Los asientos se evalúan en combinación de servicio (ELS), no con acciones mayoradas.", gray=True, size=9)
    doc.add_paragraph()
    _add_styled_table(doc, c["df"], '3.2 Estratigrafía del perfil geotécnico')
    _body(doc, f"Los módulos E deben ser coherentes con la condición de análisis adoptada. Si proceden de "
               f"ensayo edométrico (E_oed), conviene convertirlos al módulo de Young mediante "
               f"E = E_oed·(1+ν)(1−2ν)/(1−ν). El coeficiente de Poisson se limita a ν ≤ {NU_MAX} en todos los "
               f"métodos.", gray=True, size=9)
    doc.add_page_break()

    # ───────────────── 4. HIPÓTESIS Y MODELO ─────────────────
    doc.add_heading('4. Hipótesis y modelo de cálculo', level=1)
    _h2(doc, '4.1 Base común: elasticidad lineal')
    _body(doc, "Los tres métodos asumen un terreno elástico lineal, isótropo, caracterizado por (E, ν) en cada "
               "estrato. El asiento resulta de integrar la deformación vertical en profundidad hasta z_max.")
    _h2(doc, '4.2 Hipótesis de cada método')
    _bullets(doc, [
        "Steinbrenner: solución cerrada del desplazamiento del semiespacio con los factores φ₁ y φ₂; asiento por "
        "estrato como diferencia entre techo y base, con el (E, ν) del estrato.",
        "Ecuación elástica: integración numérica (Gauss, 3 puntos por subcapa) de "
        "Δεz = [Δσz − ν(Δσx+Δσy)]/E, con "
        + ("la suma de tensiones horizontales exacta para el ν de cada estrato." if consistente
           else "las tensiones horizontales clásicas de Holl (ν = 0,5)."),
        f"MEF 3D (OpenSees): continuo tridimensional con hexaedros bbarBrick, cuarto de dominio por doble "
        f"simetría y base rígida a {c['factor_prof']:.2f}·z_max.",
    ])
    doc.add_page_break()

    # ───────────────── 5. PROFUNDIDAD DE INFLUENCIA ─────────────────
    doc.add_heading('5. Profundidad de influencia (EC7)', level=1)
    _body(doc, "La profundidad de cálculo se acota mediante el criterio del Eurocódigo 7: se considera despreciable "
               "la contribución al asiento por debajo de la cota z_i en la que el incremento de tensión vertical cae "
               "por debajo del 20 % de la tensión efectiva geoestática:")
    _body(doc, "Δσz(z_i) ≤ 0,20 · σ′v0(z_i)", bold=True)
    _body(doc, f"La tensión efectiva σ′v0 incluye el terreno situado por encima de la base (D = {D:.2f} m) y se "
               f"calcula con el peso específico natural sobre el nivel freático y el peso sumergido (γsat − γw) "
               f"por debajo. Resulta z_i = {zi:.2f} m bajo la base, adoptándose z_max = {z_max:.2f} m.")
    if not c["zi_ok"]:
        _nota(doc, "El criterio del 20 % no se alcanza dentro de la estratigrafía definida: z_i se ha limitado al "
                   "espesor total del perfil. Conviene prolongar la estratigrafía.")
    doc.add_page_break()

    # ───────────────── 6. MÉTODO 1 — STEINBRENNER ─────────────────
    doc.add_heading('6. Método 1 — Steinbrenner', level=1)
    _body(doc, "Solución cerrada del asiento con los factores geométricos φ₁ y φ₂ = (m/2π)·arctan[n/(m√(1+m²+n²))]. "
               "El asiento de cada estrato se obtiene como diferencia entre el valor en su techo y en su base.")
    _add_styled_table(doc, c["df_st"], '6.1 Cálculos intermedios y asiento por estrato')
    tot1 = doc.add_paragraph(); tot1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = tot1.add_run(f'Asiento total Steinbrenner: {s_st:.3f} mm'); r.bold = True; r.font.color.rgb = RGBColor(23, 54, 93)
    doc.add_page_break()

    # ───────────────── 7. MÉTODO 2 — EC. ELÁSTICA ─────────────────
    doc.add_heading('7. Método 2 — Integración de la ecuación elástica', level=1)
    if consistente:
        _body(doc, "Integración numérica de la deformación vertical en subcapas (cuadratura de Gauss de 3 puntos). "
                   "Δσz procede de Holl y la suma Δσx+Δσy se obtiene de forma exacta para el ν de cada estrato "
                   "(θ − σz). Al compartir teoría con Steinbrenner, su coincidencia verifica la integración, no la "
                   "hipótesis elástica.")
    else:
        _body(doc, "Integración numérica de la deformación vertical en subcapas (cuadratura de Gauss de 3 puntos) con "
                   "las tensiones clásicas de Holl. Sus tensiones horizontales corresponden a ν = 0,5, por lo que el "
                   "asiento resulta inferior al de Steinbrenner cuando ν < 0,5.")
    _add_styled_table(doc, c["df_ec"], f'7.1 Tensiones medias y asiento por estrato (dz = {c["dz_sub"]} m)')
    tot2 = doc.add_paragraph(); tot2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = tot2.add_run(f'Asiento total Ec. Elástica: {s_ec:.3f} mm'); r.bold = True; r.font.color.rgb = RGBColor(33, 115, 70)
    doc.add_page_break()

    # ───────────────── 8. MÉTODO 3 — MEF 3D ─────────────────
    doc.add_heading('8. Método 3 — MEF 3D (OpenSees)', level=1)
    _h2(doc, '8.1 Descripción del modelo')
    _body(doc, f"El terreno se discretiza como un medio continuo tridimensional con hexaedros de 8 nodos con "
               f"formulación B-bar (bbarBrick), que evita el bloqueo volumétrico con ν elevado, y material "
               f"ElasticIsotropic por estrato. Se modela un cuarto del dominio por doble simetría, con una extensión "
               f"lateral de {c['factor_dominio']:.0f}·B/2 × {c['factor_dominio']:.0f}·L/2. La malla es uniforme "
               f"(tamaño {c['mesh']:.2f} m) bajo la zapata"
               + (f" y crece en progresión geométrica de razón {c['ratio']:.2f} hacia el contorno." if c['ratio'] > 1
                  else " y en el resto del dominio."))
    _h2(doc, '8.2 Condiciones de contorno y base rígida')
    _body(doc, f"En los planos de simetría y en las fronteras laterales se impide el desplazamiento normal dejando "
               f"libre el vertical. La base rígida (tres grados de libertad empotrados) se sitúa a "
               f"{c['factor_prof']*z_max:.2f} m bajo la base de la zapata ({c['factor_prof']:.2f}·z_max).")
    if mef_ok and c["factor_prof"] > 1.0:
        _body(doc, f"El asiento de comparación es s(0) − s(z_max) = {s_mef:.3f} mm, coherente con los métodos "
                   f"analíticos. El asiento total en superficie, que incluye la deformación entre z_max y la base "
                   f"rígida, es de {info['s_superficie_mm']:.3f} mm.")
        if info["extension"]:
            _nota(doc, "El perfil definido no alcanza la base rígida: el último estrato se ha prolongado con sus "
                       "mismas propiedades hasta ella.")
    if mef_ok:
        _h2(doc, '8.3 Reparto de cargas y tamaño del modelo')
        _body(doc, f"La presión se reparte en fuerzas nodales mediante áreas tributarias. Resultante aplicada: "
                   f"{info['resultante_kN']:.3f} kN frente a p·B·L/4 = {info['resultante_teorica_kN']:.3f} kN. "
                   f"Modelo: {info['n_nodos']:,} nodos y {info['n_elementos']:,} elementos; solver {info['solver']}.")
        _add_styled_table(doc, c["df_mef"], '8.4 Asiento por estrato (eje central)')
        tot3 = doc.add_paragraph(); tot3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = tot3.add_run(f'Asiento MEF 3D (0 – z_max): {s_mef:.3f} mm'); r.bold = True; r.font.color.rgb = RGBColor(192, 0, 0)
    else:
        _nota(doc, f"El cálculo MEF no ha podido completarse: {c['mef_error'].rstrip('.')}. La comprobación de servicio se "
                   f"realiza con los métodos analíticos.")
    doc.add_page_break()

    # ───────────────── 9. RESULTADOS Y COMPARATIVA ─────────────────
    doc.add_heading('9. Resultados y comparativa', level=1)
    _h2(doc, '9.1 Asientos por método (0 – z_max)')
    tr = doc.add_table(rows=1, cols=3); tr.style = 'Light Shading Accent 1'; tr.alignment = WD_TABLE_ALIGNMENT.CENTER
    cc = tr.rows[0].cells
    for cell_, (nom, val, col) in zip(cc, [('Steinbrenner', s_st, RGBColor(23, 54, 93)),
                                            ('Ec. Elástica', s_ec, RGBColor(33, 115, 70)),
                                            ('OpenSees 3D', s_mef, RGBColor(192, 0, 0))]):
        cell_.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        pp = cell_.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ra = pp.add_run(nom + '\n'); ra.font.size = Pt(11); ra.bold = True; ra.font.color.rgb = RGBColor(89, 89, 89)
        rb = pp.add_run(f'{val:.3f} mm' if np.isfinite(val) else 'No disponible')
        rb.font.size = Pt(16); rb.bold = True; rb.font.color.rgb = col
    doc.add_paragraph()
    _h2(doc, '9.2 Comparativa por estrato')
    _add_styled_table(doc, tabla_comparativa(c).drop(columns=["z Techo [m]", "z Base [m]"]), '')
    _h2(doc, '9.3 Interpretación')
    texto = f"La dispersión entre los métodos disponibles es del {dispersion:.1f} %. "
    if consistente:
        texto += ("Steinbrenner y la integración elástica resuelven el mismo problema (semiespacio elástico con el "
                  "(E, ν) de cada estrato), por lo que deben coincidir: su acuerdo verifica la implementación, no "
                  "la validez de la hipótesis. ")
    else:
        texto += ("La integración con tensiones de Holl queda por debajo de Steinbrenner porque sus tensiones "
                  "horizontales corresponden a ν = 0,5. ")
    if mef_ok:
        dif = (s_mef - s_st) / max(abs(s_st), 1e-9) * 100.0
        texto += (f"El MEF 3D difiere de Steinbrenner un {dif:+.1f} %. La diferencia es esperable: el MEF resuelve "
                  f"un estrato sobre base rígida, con dominio lateral finito y con la interacción real entre "
                  f"estratos de distinta rigidez, mientras que Steinbrenner integra el campo del semiespacio.")
    _body(doc, texto)
    doc.add_page_break()

    # ───────────────── 10. VERIFICACIÓN ─────────────────
    doc.add_heading('10. Verificación del programa', level=1)
    _body(doc, "El motor de cálculo dispone de una batería de pruebas automatizadas que contrasta: los factores de "
               "Steinbrenner con la integración numérica de la solución de Boussinesq; las tensiones de Holl con la "
               "integración directa de Boussinesq; la igualdad entre Steinbrenner y la integración elástica "
               "consistente; el equilibrio de cargas del MEF; la independencia del resultado respecto a la "
               "graduación de la malla, y la detección de fallos del solver.")
    doc.add_page_break()

    # ───────────────── 11. COMPROBACIÓN ELS ─────────────────
    doc.add_heading('11. Comprobación en Estado Límite de Servicio', level=1)
    _body(doc, "Se compara el asiento de comprobación con el asiento general admisible adoptado en función del tipo "
               "de edificio y de la naturaleza del terreno (tabla de asientos generales admisibles). "
               f"Asiento de comprobación: {origen}.")
    te = doc.add_table(rows=2, cols=4); te.style = 'Light Shading Accent 1'; te.alignment = WD_TABLE_ALIGNMENT.CENTER
    els_rows = [('Comprobación', 'Obtenido', 'Admisible', 'Veredicto'),
                ('Asiento total', f'{s_gob:.2f} mm', f'{LIM_ELS:.0f} mm', 'CUMPLE' if cumple_els else 'REVISAR')]
    for i, rowv in enumerate(els_rows):
        cs = te.rows[i].cells
        for j, v in enumerate(rowv):
            cs[j].text = v; cs[j].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for par in cs[j].paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for rr in par.runs:
                    rr.font.size = Pt(9)
                    if i == 0: rr.font.bold = True; rr.font.color.rgb = RGBColor(23, 54, 93)
    doc.add_paragraph()
    _body(doc, f"El asiento obtenido ({s_gob:.2f} mm) supone un aprovechamiento del {s_gob/LIM_ELS*100:.0f} % del "
               f"asiento admisible adoptado ({LIM_ELS:.0f} mm).", gray=True, size=9)
    if c["cohesivos"]:
        _nota(doc, "Estratos cohesivos dentro de z_max (" + ", ".join(c["cohesivos"]) + "): el asiento calculado es "
                   "elástico. Si los módulos no son drenados, falta sumar el asiento de consolidación antes de "
                   "comparar con el asiento total admisible.")
    doc.add_page_break()

    # ───────────────── 12. ZONA DE INFLUENCIA DE LA CARGA ─────────────────
    doc.add_heading('12. Zona de influencia de la carga de cimentación', level=1)
    pf = doc.add_paragraph(); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.add_run().add_picture(fig_bulbo_bytes, width=Cm(13))
    nota = doc.add_paragraph(f'Evolución de tensiones bajo el centro de la zapata (p={p:.1f} kPa, B={B:.2f} m, '
                             f'L={L:.2f} m). Criterio EC7: 0,20·σ′v0.')
    nota.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nota.runs[0].font.size = Pt(9); nota.runs[0].font.italic = True; nota.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════
# CONFIGURACIÓN E INTERFAZ STREAMLIT
# ══════════════════════════════════════════════════════════════════════════
def reset_calculo():
    # Un cambio en los datos invalida el cálculo Y el informe (que se basa en él)
    st.session_state.calculo_realizado = False
    if st.session_state.get("informe_generado", False):
        st.session_state.informe_stale = True
    st.session_state.informe_generado = False
    st.session_state.word_buf = None


@st.cache_data(show_spinner=False)
def _zi_cacheado(p, B, L, df, NF, D, gamma_D, gamma_sat_D):
    return z_influencia_ec7(p, B, L, df, NF, D, gamma_D, gamma_sat_D)


for _k, _v in [("calculo_realizado", False), ("informe_generado", False),
               ("informe_stale", False), ("word_buf", None), ("calc", None)]:
    if _k not in st.session_state:
        st.session_state[_k] = _v

if 'df_terreno' not in st.session_state:
    st.session_state.df_terreno = pd.DataFrame({
        "Descripción":           ["Arcilla 1", "Arcilla 2", "Grava"],
        "Tipo":                  ["Cohesivo", "Cohesivo", "Granular"],
        "Espesor (m)":           [1.5, 3.0, 5.0],
        "E (kPa)":               [10000.0, 15000.0, 40000.0],
        "nu":                    [0.30, 0.45, 0.25],
        "Peso Esp. (kN/m³)":     [18.0, 19.0, 21.0],
        "Peso Esp. Sat (kN/m³)": [20.0, 20.0, 22.0],
    })

st.set_page_config(page_title="Cálculo de Asientos · MEF 3D", layout="wide", page_icon="🏗️")
st.sidebar.title("Navegación")
modo = st.sidebar.radio("Vista:", ["🧮 Panel de Cálculo", "📋 Modelo Steinbrenner", "📋 Modelo Elástico",
                                   "🌐 Modelo OpenSees", "📉 Bulbo de Presiones", "📐 Asientos Admisibles",
                                   "📖 Fundamento Teórico"])

# ── DATOS DE ENTRADA ──
st.sidebar.markdown("---")
st.sidebar.header("📥 Datos de Entrada")
B  = st.sidebar.number_input("Ancho (B) [m]", min_value=0.1, value=2.0, step=0.1, on_change=reset_calculo)
L  = st.sidebar.number_input("Longitud (L) [m]", min_value=0.1, value=3.0, step=0.1, on_change=reset_calculo)
p  = st.sidebar.number_input("Presión neta (p) [kPa]", min_value=1.0, value=150.0, step=10.0, on_change=reset_calculo)
D  = st.sidebar.number_input("Profundidad de cimentación (D) [m]", min_value=0.0, value=0.0, step=0.1,
                             on_change=reset_calculo, help="Desde la superficie del terreno hasta la base de la zapata.")
NF = st.sidebar.number_input("Nivel freático, desde superficie [m]", min_value=0.0, value=100.0, step=0.5,
                             on_change=reset_calculo, help="Un valor ≥ 100 m equivale a «sin nivel freático».")
with st.sidebar.expander("Terreno por encima de la base (0 ≤ prof. ≤ D)"):
    gamma_D = st.number_input("γ [kN/m³]", min_value=10.0, value=18.0, step=0.5, on_change=reset_calculo)
    gamma_sat_D = st.number_input("γsat [kN/m³]", min_value=10.0, value=20.0, step=0.5, on_change=reset_calculo)

if L < B:
    B, L = L, B
    st.sidebar.warning("⚠️ L<B: valores intercambiados.")

df_val, errores, avisos = validar_estratigrafia(st.session_state.df_terreno)
datos_ok = not errores
if datos_ok:
    espesor_total = max(float(df_val["Espesor (m)"].sum()), 0.1)
    zi, zi_ok = _zi_cacheado(p, B, L, df_val, NF, D, gamma_D, gamma_sat_D)
else:
    espesor_total, zi, zi_ok = 0.1, 0.1, False
    st.sidebar.error("❌ La estratigrafía tiene errores (ver panel de cálculo).")

# ── PROFUNDIDAD DE CÁLCULO ──
st.sidebar.markdown("---")
st.sidebar.subheader("📐 Profundidad de Cálculo")
if datos_ok and not zi_ok:
    st.sidebar.warning("El criterio EC7 no se alcanza dentro del perfil: z_i se limita al espesor total. "
                       "Conviene prolongar la estratigrafía.")
z_auto = float(max(0.1, min(round(zi, 2), espesor_total)))
usar_zi = st.sidebar.checkbox("Usar z_i (EC7) como profundidad de corte", value=True, key="z_max_auto",
                              on_change=reset_calculo)
if usar_zi:
    z_max_user = z_auto
    st.sidebar.markdown(f"**z_max = z_i = {z_max_user:.2f} m** bajo la base")
else:
    if "z_max_manual" not in st.session_state:
        st.session_state.z_max_manual = z_auto
    st.session_state.z_max_manual = float(min(max(st.session_state.z_max_manual, 0.1), espesor_total))
    z_max_user = st.sidebar.number_input("Profundidad de corte (z_max) [m]", min_value=0.1,
                                         max_value=espesor_total, step=0.1, key="z_max_manual",
                                         on_change=reset_calculo)
    st.sidebar.caption(f"Referencia EC7: z_i = {zi:.2f} m")

# ── ELS ──
st.sidebar.markdown("---")
st.sidebar.subheader("✅ Comprobación de Servicio (ELS)")
asiento_adm = st.sidebar.number_input(
    "Asiento admisible [mm]", min_value=1.0, value=25.0, step=1.0, on_change=reset_calculo,
    help="Consulta la pestaña «📐 Asientos Admisibles» para elegirlo según edificio y terreno.")
criterio = st.sidebar.selectbox("Asiento de comprobación", CRITERIOS_GOBERNANTE, index=0, on_change=reset_calculo,
                                help="Si se elige el MEF y no puede resolverse, se usa el máximo de los analíticos.")
factor_rigidez = st.sidebar.number_input(
    "Factor de rigidez de la zapata (s_rígida / s_centro)", min_value=0.5, max_value=1.0, value=1.0, step=0.01,
    on_change=reset_calculo,
    help="1,0 = centro de zapata flexible (lado seguro). Para zapata rígida la bibliografía propone valores "
         "del orden de 0,8–0,93 según la fuente; justifica el valor adoptado.")

# ── MÉTODO 2 ──
st.sidebar.markdown("---")
st.sidebar.subheader("🔧 Ecuación Elástica")
dz_sub = st.sidebar.select_slider("Tamaño de subcapa (dz) [m]", options=[2.0, 1.0, 0.5, 0.25, 0.10, 0.05],
                                  value=0.25, on_change=reset_calculo)
_modo_lbl = st.sidebar.radio("Tensiones horizontales",
                             ["Consistentes con ν (exactas)", "Holl clásicas (ν = 0,5)"],
                             on_change=reset_calculo,
                             help="Las expresiones de Holl para σx, σy suponen ν = 0,5 e infravaloran el asiento "
                                  "cuando ν es menor.")
modo_tension = TENSION_CONSISTENTE if _modo_lbl.startswith("Consistentes") else TENSION_HOLL

# ── MEF 3D ──
st.sidebar.markdown("---")
st.sidebar.subheader("⚙️ Parámetros OpenSees MEF 3D")
if not OPENSEES_DISPONIBLE:
    st.sidebar.warning("⚠️ OpenSeesPy no detectado: la comprobación usará los métodos analíticos.")
factor_dominio = st.sidebar.slider("Extensión del dominio (×B/2, ×L/2)", min_value=3.0, max_value=8.0,
                                   value=5.0, step=1.0, on_change=reset_calculo)
mesh_3d = st.sidebar.number_input("Tamaño de malla bajo la zapata [m]", min_value=0.1, max_value=1.0, value=0.5,
                                  step=0.05, on_change=reset_calculo)
ratio_malla = st.sidebar.select_slider("Crecimiento de la malla hacia el contorno",
                                       options=[1.0, 1.1, 1.2, 1.3, 1.5], value=1.2, on_change=reset_calculo,
                                       help="1,0 = malla uniforme. Valores > 1 engrosan la malla fuera de la zapata.")
factor_prof = st.sidebar.slider("Base rígida del MEF (× z_max)", min_value=1.0, max_value=2.0, value=1.0,
                                step=0.25, on_change=reset_calculo,
                                help="1,0 = base rígida en z_max (comparable con los analíticos). Con valores > 1 "
                                     "se informa además el asiento total en superficie.")
if factor_dominio < 4:
    st.sidebar.warning("Con factor < 4 las fronteras laterales rigidizan el modelo y subestiman el asiento (≥5 recomendado).")

if datos_ok:
    nodos_totales, elementos_totales = estimar_malla_3d(B, L, df_val, z_max_user, mesh_3d,
                                                       factor_dominio, factor_prof, ratio_malla)
    if nodos_totales > MAX_NODOS_MEF:
        st.sidebar.error(f"🔴 **Malla excesiva:** {nodos_totales:,} nodos (límite {MAX_NODOS_MEF:,}). "
                         "El MEF no se ejecutará.")
    elif nodos_totales > 10000:
        st.sidebar.warning(f"🟡 **Malla densa:** {nodos_totales:,} nodos · {elementos_totales:,} elem. "
                           "Puede tardar del orden de un minuto.")
    else:
        st.sidebar.success(f"🟢 **Malla ligera:** {nodos_totales:,} nodos · {elementos_totales:,} elem.")

# ── CÁLCULO ──
st.sidebar.markdown("---")
if st.sidebar.button("🚀 Calcular", type="primary", width="stretch", disabled=not datos_ok):
    with st.spinner("Calculando…"):
        tot_st, df_st = calcular_steinbrenner(p, B, L, df_val, z_max_user)
        tot_ec, df_ec = calcular_ec68(p, B, L, df_val, z_max_user, dz_sub, modo_tension)
        try:
            tot_mef, df_mef, mef_info = calcular_opensees_3d(p, B, L, df_val, z_max_user, mesh_3d,
                                                             factor_dominio, factor_prof, ratio_malla)
            mef_error = None
        except MEFError as exc:
            tot_mef, mef_info, mef_error = float("nan"), None, str(exc)
            df_mef = pd.DataFrame({"Capa": df_st["Capa"], "Δs [mm]": np.nan})
    st.session_state.calc = dict(
        B=B, L=L, p=p, D=D, NF=NF, gamma_D=gamma_D, gamma_sat_D=gamma_sat_D, df=df_val,
        zi=zi, zi_ok=zi_ok, z_max=z_max_user, s_adm=asiento_adm, criterio=criterio,
        factor_rigidez=factor_rigidez, dz_sub=dz_sub, modo_tension=modo_tension,
        factor_dominio=factor_dominio, mesh=mesh_3d, ratio=ratio_malla, factor_prof=factor_prof,
        tot_st=tot_st, df_st=df_st, tot_ec=tot_ec, df_ec=df_ec,
        tot_mef=tot_mef, df_mef=df_mef, mef_info=mef_info, mef_error=mef_error,
        cohesivos=estratos_cohesivos(df_val, z_max_user),
    )
    st.session_state.calculo_realizado = True
    if st.session_state.get("informe_generado", False):
        st.session_state.informe_stale = True
    st.session_state.informe_generado = False
    st.session_state.word_buf = None

calc = st.session_state.calc if st.session_state.calculo_realizado else None

# ── DOCUMENTACIÓN ──
st.sidebar.markdown("---")
st.sidebar.subheader("🗎 Documentación")
with st.sidebar.expander("🗂️ Identificación del informe"):
    meta_obra = st.text_input("Obra / Proyecto", value="")
    meta_peti = st.text_input("Peticionario", value="")
    meta_ref  = st.text_input("Referencia", value="GEO-XXXX-ASN-01")
    meta_autor = st.text_input("Autor", value="Dpto. Geotecnia")
meta_informe = {"obra": meta_obra, "peticionario": meta_peti, "referencia": meta_ref, "autor": meta_autor}

if st.session_state.get("informe_stale", False):
    st.sidebar.warning("⚠️ Los datos han cambiado. **Recalcula** y vuelve a **generar el informe** "
                       "para que el documento corresponda a los resultados actuales.")

if calc is None:
    if not st.session_state.get("informe_stale", False):
        st.sidebar.info("Ejecuta el cálculo para poder generar el informe.")
else:
    if st.sidebar.button("📝 Generar informe", width="stretch"):
        with st.spinner("Generando memoria de cálculo…"):
            fig_b = figura_bulbo(calc, float(calc["df"]["Espesor (m)"].sum()))
            fig_bulbo_bytes = _fig_bytes(fig_b); plt.close(fig_b)
            st.session_state.word_buf = generar_word(calc, fig_bulbo_bytes, meta=meta_informe)
        st.session_state.informe_generado = True
        st.session_state.informe_stale = False
        st.sidebar.success("✅ Informe generado.")

if st.session_state.get("informe_generado", False) and st.session_state.get("word_buf") is not None:
    st.sidebar.download_button("⬇️ Descargar informe Word", data=st.session_state.word_buf,
                               file_name="informe_comparativo.docx", width="stretch")

# ══════════════════════════════════════════════════════════════════════════
# ÁREA PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════
st.title("🏗️ Cálculo de Asientos de cimentaciones rectangulares")
st.markdown("**Herramienta académica** · los resultados deben ser revisados por un técnico competente")
st.markdown("Según directiva ITQ404")

if modo == "🧮 Panel de Cálculo":
    st.header("1. Estratigrafía del Terreno")
    st.caption("Profundidades medidas desde la base de la zapata. «Tipo» se usa para los avisos de ELS.")
    df_edit = st.data_editor(
        st.session_state.df_terreno, num_rows="dynamic", width="stretch",
        column_config={"Tipo": st.column_config.SelectboxColumn("Tipo", options=TIPOS_SUELO)})
    if not df_edit.equals(st.session_state.df_terreno):
        st.session_state.df_terreno = df_edit
        reset_calculo()
        st.rerun()
    for e in errores:
        st.error(e)
    for a in avisos:
        st.warning(a)

    st.markdown("---")
    st.header("2. Resultados y Comparativa")
    if calc is None:
        st.info("👈 Pulsa Calcular.")
    else:
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("🔵 Steinbrenner", _fmt_mm(calc["tot_st"]))
        c2.metric("🟢 Ec. Elástica", _fmt_mm(calc["tot_ec"]))
        c3.metric("🔴 OpenSees 3D", _fmt_mm(calc["tot_mef"]))
        if np.isfinite(calc["tot_mef"]):
            dif = (calc["tot_mef"] - calc["tot_st"]) * 1000
            pct = dif / max(abs(calc["tot_st"]) * 1000, 1e-9) * 100
            c4.metric("📊 MEF − Steinbrenner", f"{dif:+.3f} mm", f"{pct:+.1f} %", delta_color="off")
        else:
            st.error(f"❌ MEF 3D no disponible: {calc['mef_error']}")

        s_gob, origen = asiento_gobernante(calc["tot_st"], calc["tot_ec"], calc["tot_mef"],
                                           calc["criterio"], calc["factor_rigidez"])
        aprov = s_gob / calc["s_adm"] * 100
        msg = (f"Asiento de comprobación ({origen}) {s_gob:.2f} mm "
               f"{'≤' if s_gob <= calc['s_adm'] else '>'} admisible {calc['s_adm']:.0f} mm · "
               f"aprovechamiento {aprov:.0f} %.")
        if s_gob <= calc["s_adm"]:
            st.success("✅ **CUMPLE** — " + msg)
        else:
            st.error("⚠️ **REVISAR** — " + msg)
        if calc["cohesivos"]:
            st.warning("⚠️ Hay estratos cohesivos dentro de z_max (" + ", ".join(calc["cohesivos"]) + "). "
                       "El asiento calculado es elástico: si los módulos no son drenados, falta la consolidación "
                       "antes de compararlo con el asiento total admisible.")
        st.dataframe(tabla_comparativa(calc), width="stretch", hide_index=True)

# ══════════════════════════════════════════════════
# VISTA: ASIENTOS ADMISIBLES
# ══════════════════════════════════════════════════
elif modo == "📐 Asientos Admisibles":
    st.header("📐 Asientos Generales Admisibles")
    st.markdown("Tabla de referencia para fijar el **asiento admisible** en la barra lateral, según el tipo de "
                "edificio y la naturaleza del terreno. El valor elegido se emplea en la comprobación ELS y en el informe.")
    df_adm = pd.DataFrame({
        "Características del edificio": [
            "Obras de carácter monumental",
            "Edificios con estructura de H.A. de gran rigidez",
            "Edificios con estructura de H.A. de pequeña rigidez · Estructuras metálicas hiperestáticas · Edificios con muros de fábrica",
            "Estructuras metálicas isostáticas · Estructuras de madera · Estructuras provisionales",
        ],
        "Terreno sin cohesión [mm]": ["12", "35", "50", ">50 (con comprobación)"],
        "Terreno coherente [mm]": ["25", "50", "75", ">75 (con comprobación)"],
    })
    st.dataframe(df_adm, width="stretch", hide_index=True)
    st.info(f"**Asiento admisible fijado actualmente:** {asiento_adm:.0f} mm  "
            f"(se edita en la barra lateral, apartado «Comprobación de Servicio»).")
    st.markdown("**Notas:**")
    st.markdown(
        "- *Sin cohesión* = terrenos granulares (arenas, gravas); *coherentes* = terrenos cohesivos (arcillas, limos). "
        "Los límites admisibles son mayores en terrenos coherentes.\n"
        "- Son asientos **totales**: en terrenos cohesivos el asiento elástico calculado aquí no incluye la "
        "consolidación, salvo que se hayan empleado módulos drenados.\n"
        "- La última fila (*«con comprobación»*) no es un límite cerrado: indica que se admiten asientos mayores "
        "siempre que se justifique que la estructura los tolera.\n"
        "- En perfiles multicapa mixtos, la elección del tipo de terreno y de edificio es criterio del proyectista.")
    st.caption("Fuente: tabla de asientos generales admisibles (Jiménez Salas), de uso habitual en la práctica geotécnica.")

# ══════════════════════════════════════════════════
# VISTA: DETALLE STEINBRENNER
# ══════════════════════════════════════════════════
elif modo == "📋 Modelo Steinbrenner":
    st.header("📋 Detalle Método Steinbrenner")
    st.markdown(r"Cálculo capa a capa con las funciones de influencia $\phi_1$ y $\phi_2$ "
                r"(4 cuadrantes de $B/2 \times L/2$).")
    if calc is None:
        st.warning("⚠️ Calcula primero en el panel izquierdo.")
    else:
        df_st = calc["df_st"]
        st.markdown("##### 🔼 Valores en el Techo")
        st.dataframe(df_st[["Capa", "z Techo [m]", "m_techo", "φ1_techo", "φ2_techo", "s_techo [mm]"]],
                     width="stretch", hide_index=True)
        st.markdown("##### 🔽 Valores en la Base")
        st.dataframe(df_st[["Capa", "z Base [m]", "m_base", "φ1_base", "φ2_base", "s_base [mm]"]],
                     width="stretch", hide_index=True)
        st.markdown("##### 📊 Asiento por estrato")
        st.dataframe(df_st[["Capa", "Δs [mm]"]], width="stretch", hide_index=True)
        st.metric("🔵 Asiento Total Steinbrenner", _fmt_mm(calc["tot_st"]))

# ══════════════════════════════════════════════════
# VISTA: DETALLE EC ELÁSTICO
# ══════════════════════════════════════════════════
elif modo == "📋 Modelo Elástico":
    st.header("📋 Detalle Método Ecuación Elástica")
    st.latex(r"s = \sum_i \int_{z_{t,i}}^{z_{b,i}} \frac{\Delta\sigma_z - \nu_i(\Delta\sigma_x+\Delta\sigma_y)}{E_i}\,dz")
    if calc is None:
        st.warning("⚠️ Calcula primero.")
    else:
        df_ec = calc["df_ec"]
        tipo = ("suma horizontal exacta para el ν de cada estrato" if calc["modo_tension"] == TENSION_CONSISTENTE
                else "tensiones horizontales de Holl (ν = 0,5)")
        st.caption(f"Subcapas de **{calc['dz_sub']} m**, cuadratura de Gauss de 3 puntos, {tipo}. "
                   "Los valores medios son promedios ponderados en el espesor del estrato.")
        st.markdown("##### ⚡ Tensiones medias por capa")
        st.dataframe(df_ec[["Capa", "Sub-capas", "Δσz med [kPa]", "Δ(σx+σy) med [kPa]"]],
                     width="stretch", hide_index=True)
        st.markdown("##### 📐 Deformación unitaria media y asiento")
        st.dataframe(df_ec[["Capa", "h_ef [m]", "Sub-capas", "Δεz med [-]", "Δs [mm]"]],
                     width="stretch", hide_index=True)
        st.metric("🟢 Asiento Total Ec. Elástica", _fmt_mm(calc["tot_ec"]))

# ══════════════════════════════════════════════════
# VISTA: MODELO OPENSEES
# ══════════════════════════════════════════════════
elif modo == "🌐 Modelo OpenSees":
    st.header("🌐 Modelo de Elementos Finitos (OpenSees 3D)")
    st.markdown("El terreno se resuelve como un **medio continuo tridimensional** discretizado por el Método de los "
                "Elementos Finitos: el asiento de la zapata rectangular sale directamente del cálculo.")
    st.subheader("1. Concepto y Datos de Entrada")
    st.markdown("* **Materiales:** a cada estrato se le asigna un modelo `ElasticIsotropic` ($E$, $\\nu$).\n"
                "* **Elementos:** hexaedros de 8 nodos con formulación B-bar (`bbarBrick`), que evita el bloqueo "
                "volumétrico con $\\nu$ elevado, en un **cuarto del dominio** por doble simetría.\n"
                "* **Malla:** uniforme bajo la zapata, con el borde cargado siempre en un nodo, y en progresión "
                "geométrica hacia el contorno si el crecimiento es > 1.\n"
                "* **Carga:** presión repartida en fuerzas nodales por **áreas tributarias**, con resultante exacta "
                "$p\\cdot B\\cdot L/4$.")
    st.subheader("2. Condiciones de Contorno")
    col1, col2 = st.columns(2)
    with col1:
        st.info("📏 Simetría y campo lejano")
        st.markdown(r"Planos de simetría ($x=0$, $y=0$) y fronteras laterales: desplazamiento normal impedido, "
                    r"vertical libre.")
    with col2:
        st.info("⬇️ Base")
        st.markdown(rf"Base rígida (3 GDL empotrados) a **{factor_prof:.2f}·z_max**. El asiento de comparación es "
                    r"siempre $s(0)-s(z_{max})$, coherente con los métodos analíticos.")
    if factor_prof > 1.0:
        st.info("Con la base por debajo de z_max se informa además el asiento total en superficie. Si el perfil no "
                "llega a la base rígida, el último estrato se prolonga con sus mismas propiedades.")
    st.markdown("---")
    if calc is None:
        st.warning("⚠️ Ejecuta el cálculo en el panel izquierdo para visualizar la extracción de asientos.")
    elif calc["mef_info"] is None:
        st.error(f"❌ El MEF no ha podido resolverse: {calc['mef_error']}")
    else:
        info = calc["mef_info"]
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Nodos", f"{info['n_nodos']:,}")
        m2.metric("Elementos", f"{info['n_elementos']:,}")
        m3.metric("Tiempo", f"{info['t_calculo_s']:.1f} s")
        m4.metric("Resultante", f"{info['resultante_kN']:.2f} kN",
                  f"teórica {info['resultante_teorica_kN']:.2f} kN", delta_color="off")
        st.caption(f"Solver lineal: {info['solver']}")
        st.markdown("##### 📍 Asiento por estrato (eje central)")
        st.dataframe(calc["df_mef"][["Capa", "Δs [mm]"]], width="stretch", hide_index=True)
        st.metric("🔴 Asiento OpenSees 3D (0 – z_max)", _fmt_mm(calc["tot_mef"]))
        if calc["factor_prof"] > 1.0:
            st.metric("Asiento total en superficie (hasta la base rígida)", f"{info['s_superficie_mm']:.3f} mm")
            if info["extension"]:
                st.warning("El perfil no alcanza la base rígida: el último estrato se ha prolongado.")

# ══════════════════════════════════════════════════
# VISTA: BULBO DE PRESIONES
# ══════════════════════════════════════════════════
elif modo == "📉 Bulbo de Presiones":
    st.header("Bulbo de Presiones y Zona de Influencia")
    st.markdown(r"Tensiones bajo el centro ($\times 4$ superposición $B/2 \times L/2$). "
                r"El criterio EC7 es: $\Delta\sigma_z \leq 0.20\,\sigma'_{v0}$.")
    if not datos_ok:
        st.warning("Corrige la estratigrafía para ver el bulbo.")
    else:
        col1, col2 = st.columns([1, 3])
        with col1:
            z_gr = st.slider("Profundidad máxima [m]:", 0.1, max(espesor_total, 0.2), min(max(espesor_total, 0.2), 15.0), 0.1)
            st.markdown(f"**p:** {p} kPa · **B:** {B} m · **L:** {L} m · **D:** {D} m")
            if NF < 100.0: st.markdown(f"**NF:** {NF:.1f} m desde superficie")
            st.metric("📐 z_i (EC7)", f"{zi:.2f} m")
            st.markdown("---")
            st.info("Δσz es común a los dos métodos analíticos. Las curvas de Δσx y Δσy son las de Holl (ν = 0,5); "
                    "el método elástico consistente usa la suma exacta para el ν de cada estrato.")
        with col2:
            ctx = dict(p=p, B=B, L=L, df=df_val, NF=NF, D=D, gamma_D=gamma_D, gamma_sat_D=gamma_sat_D, zi=zi)
            fig = figura_bulbo(ctx, z_gr, figsize=(9, 7))
            st.pyplot(fig); plt.close(fig)

# ══════════════════════════════════════════════════
# VISTA: FUNDAMENTO TEÓRICO
# ══════════════════════════════════════════════════
elif modo == "📖 Fundamento Teórico":
    st.header("Fundamento Teórico — Comparativa de Formulaciones")
    col_a, col_b = st.columns(2)
    with col_a:
        st.subheader("🔵 Método 1 — Steinbrenner")
        st.markdown(r"Desplazamiento vertical del semiespacio bajo la esquina de un rectángulo $B' \times L'$; "
                    r"para el centro se superponen 4 cuadrantes con $B'=B/2$, $L'=L/2$:")
        st.latex(r"s(z) = 4\,\frac{p\,B'}{E}\left[(1-\nu^2)\,\phi_1 - (1-\nu-2\nu^2)\,\phi_2\right]")
        st.latex(r"\phi_1 = \frac{1}{\pi}\left[\ln\frac{\sqrt{1+m^2+n^2}+n}{\sqrt{1+m^2}} + n\ln\frac{\sqrt{1+m^2+n^2}+1}{\sqrt{n^2+m^2}}\right]")
        st.latex(r"\phi_2 = \frac{m}{2\pi}\arctan\frac{n}{m\sqrt{1+m^2+n^2}}")
        st.markdown(r"Con $n = L'/B' = L/B$ y $m = z/B' = 2z/B$. El asiento de cada estrato:")
        st.latex(r"\Delta s_i = s(z_{techo}) - s(z_{base})")
    with col_b:
        st.subheader("🟢 Método 2 — Ecuación Elástica")
        st.markdown("Integración numérica de la deformación vertical en subcapas (Gauss, 3 puntos por subcapa):")
        st.latex(r"\Delta\varepsilon_z = \frac{\Delta\sigma_z - \nu(\Delta\sigma_x+\Delta\sigma_y)}{E}")
        st.markdown(r"La suma de tensiones horizontales exacta para $\nu$ arbitrario se obtiene de la tensión "
                    r"media de Boussinesq integrada sobre el rectángulo:")
        st.latex(r"\Delta\sigma_x+\Delta\sigma_y = (1+\nu)\,\frac{p}{\pi}\arctan\frac{BL}{zR_3} - \Delta\sigma_z"
                 r"\quad\text{(por esquina)}")
        st.info("Con esta suma, el método 2 coincide con Steinbrenner: sirve de verificación numérica. "
                "Con las tensiones clásicas de Holl (ν = 0,5) el asiento queda por debajo cuando ν < 0,5.")
    st.markdown("---")
    st.subheader("🔁 Tensiones de Holl bajo la esquina")
    st.markdown(r"Superposición ×4 con $B/2 \times L/2$ para obtener el **centro** de la zapata. "
                r"$\sigma_z$ es válida para cualquier $\nu$; $\sigma_x$ y $\sigma_y$ corresponden a $\nu = 0{,}5$.")
    st.latex(r"\sigma_z = \frac{p}{2\pi}\left[\arctan\frac{BL}{zR_3} + BL\left(\frac{1}{R_1^2}+\frac{1}{R_2^2}\right)\frac{z}{R_3}\right]")
    st.latex(r"\sigma_x = \frac{p}{2\pi}\left[\arctan\frac{BL}{zR_3} - \frac{BLz}{R_1^2 R_3}\right]")
    st.latex(r"\sigma_y = \frac{p}{2\pi}\left[\arctan\frac{BL}{zR_3} - \frac{BLz}{R_2^2 R_3}\right]")
    st.latex(r"R_1=\sqrt{L^2+z^2}\quad R_2=\sqrt{B^2+z^2}\quad R_3=\sqrt{L^2+B^2+z^2}")
    st.markdown("---")
    st.subheader("📐 Criterio de Profundidad de Influencia (UNE-EN 1997-1, 6.6.2)")
    st.latex(r"\Delta\sigma_z(z_i) \leq 0.20\,\sigma'_{v0}(z_i)")
    st.markdown(r"Con $\sigma'_{v0}$ = tensión efectiva geoestática, incluido el terreno sobre la base (D) y "
                r"considerando el nivel freático: $\gamma_i$ sobre NF y $\gamma_{sat,i} - \gamma_w$ bajo NF.")
