from __future__ import annotations

from io import BytesIO
from typing import Any

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from calculos.asientos import ResultadoCalculo
from calculos.steinbrenner import ResultadoSteinbrenner
from modelos.datos import Cimentacion


AZUL = "2164AF"
AZUL_OSCURO = "12304A"


def _borde_inferior(parrafo, color: str = AZUL) -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def _quitar_bordes(parrafo) -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def _configurar_documento(doc: Document) -> None:
    seccion = doc.sections[0]
    seccion.page_width = Mm(210)
    seccion.page_height = Mm(297)
    seccion.top_margin = Mm(25)
    seccion.bottom_margin = Mm(25)
    seccion.left_margin = Mm(30)
    seccion.right_margin = Mm(30)
    seccion.header_distance = Mm(12.5)
    seccion.footer_distance = Mm(12.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Open Sans"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal.paragraph_format.space_after = Pt(8)

    titulo = doc.styles["Title"]
    titulo_p_pr = titulo._element.get_or_add_pPr()
    titulo_bordes = titulo_p_pr.find(qn("w:pBdr"))
    if titulo_bordes is not None:
        titulo_p_pr.remove(titulo_bordes)

    for nombre, tamano, cursiva in [("Heading 1", 12, False), ("Heading 2", 10, True)]:
        estilo = doc.styles[nombre]
        estilo.font.name = "Open Sans Condensed"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.italic = cursiva
        estilo.font.color.rgb = RGBColor.from_string(AZUL)
        estilo.paragraph_format.keep_with_next = True
        estilo.paragraph_format.space_before = Pt(18 if nombre == "Heading 1" else 12)
        estilo.paragraph_format.space_after = Pt(6)

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pie.add_run("ASIENTO ELÁSTICO DE ZAPATA RECTANGULAR FLEXIBLE")
    run.font.name = "Open Sans"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(AZUL)


def _titulo(doc: Document, texto: str, nivel: int = 1) -> None:
    p = doc.add_heading(texto, level=nivel)
    if nivel == 1:
        _borde_inferior(p)


def _tabla(doc: Document, df: pd.DataFrame, decimales: int = 3) -> None:
    tabla = doc.add_table(rows=1, cols=len(df.columns))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = True
    for j, nombre in enumerate(df.columns):
        celda = tabla.rows[0].cells[j]
        celda.text = str(nombre)
        celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in celda.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
            run.font.size = Pt(8)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), AZUL)
        celda._tc.get_or_add_tcPr().append(shd)
    for _, fila in df.iterrows():
        celdas = tabla.add_row().cells
        for j, valor in enumerate(fila):
            if pd.isna(valor):
                texto = ""
            elif isinstance(valor, (float, np.floating)):
                texto = f"{valor:.{decimales}f}"
            else:
                texto = str(valor)
            celdas[j].text = texto
            celdas[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in celdas[j].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Open Sans"
                    run.font.size = Pt(7.5)
    doc.add_paragraph()


def _figura_tensiones(resultado: ResultadoCalculo) -> BytesIO:
    centro = resultado.perfiles[resultado.perfiles["Punto"] == "Centro"]
    fig, ax = plt.subplots(figsize=(6.0, 4.0), dpi=180)
    ax.plot(centro["sigma_v0_total_kPa"], centro["z_bajo_base_m"], label="σv0 total", color="#778497")
    ax.plot(centro["sigma_v0_ef_kPa"], centro["z_bajo_base_m"], label="σ'v0 efectiva", color="#17835C")
    ax.plot(centro["delta_sigma_z_kPa"], centro["z_bajo_base_m"], label="Δσz", color="#E45454")
    ax.plot(centro["sigma_ef_final_kPa"], centro["z_bajo_base_m"], label="σ'v0 + Δσz", color="#2164AF", linestyle="--")
    ax.invert_yaxis(); ax.grid(alpha=0.25); ax.legend(frameon=False, fontsize=8)
    ax.set_xlabel("Tensión vertical (kPa)"); ax.set_ylabel("Profundidad bajo la base (m)")
    fig.tight_layout()
    salida = BytesIO(); fig.savefig(salida, format="png", bbox_inches="tight"); plt.close(fig); salida.seek(0)
    return salida


def _figura_asientos(resultado: ResultadoCalculo) -> BytesIO:
    fig, ax = plt.subplots(figsize=(6.0, 4.0), dpi=180)
    for nombre, grupo in resultado.perfiles.groupby("Punto", sort=False):
        ax.plot(grupo["desplazamiento_mm"], grupo["z_bajo_base_m"], label=nombre)
    ax.invert_yaxis(); ax.grid(alpha=0.25); ax.legend(frameon=False, fontsize=7, ncol=2)
    ax.set_xlabel("Desplazamiento vertical (mm)"); ax.set_ylabel("Profundidad bajo la base (m)")
    fig.tight_layout()
    salida = BytesIO(); fig.savefig(salida, format="png", bbox_inches="tight"); plt.close(fig); salida.seek(0)
    return salida


def _figura_bulbo(resultado: ResultadoCalculo, cimentacion: Cimentacion) -> BytesIO:
    malla = resultado.malla_corta
    fig, ax = plt.subplots(figsize=(6.0, 4.3), dpi=180)
    porcentaje = 100.0 * malla["delta_sigma_kPa"] / cimentacion.q_neta_kpa
    cf = ax.contourf(malla["x_m"], malla["z_m"], porcentaje, levels=np.arange(5, 85, 5), cmap="RdYlBu_r", extend="max")
    ax.contour(malla["x_m"], malla["z_m"], porcentaje, levels=[10, 20, 40, 60, 80], colors="#12304A", linewidths=0.65)
    ax.plot([-cimentacion.b_m / 2, cimentacion.b_m / 2], [0, 0], color="#12304A", linewidth=4)
    ax.invert_yaxis(); ax.set_aspect("equal", adjustable="box")
    ax.set_xlabel("Distancia en dirección B (m)"); ax.set_ylabel("Profundidad bajo la base (m)")
    fig.colorbar(cf, ax=ax, label="Δσz/q (%)"); fig.tight_layout()
    salida = BytesIO(); fig.savefig(salida, format="png", bbox_inches="tight"); plt.close(fig); salida.seek(0)
    return salida


def _figura_steinbrenner(resultado: ResultadoSteinbrenner) -> BytesIO:
    fig, ax = plt.subplots(figsize=(6.0, 4.0), dpi=180)
    for nombre, grupo in resultado.perfiles.groupby("Punto", sort=False):
        ax.plot(
            grupo["desplazamiento_mm"], grupo["z_bajo_base_m"],
            marker="o", markersize=3.5, label=nombre,
        )
    ax.invert_yaxis(); ax.grid(alpha=0.25); ax.legend(frameon=False, fontsize=7, ncol=2)
    ax.set_xlabel("Desplazamiento Steinbrenner (mm)")
    ax.set_ylabel("Profundidad bajo la base (m)")
    fig.tight_layout()
    salida = BytesIO(); fig.savefig(salida, format="png", bbox_inches="tight"); plt.close(fig); salida.seek(0)
    return salida


def _insertar_figura(doc: Document, imagen: BytesIO, pie: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(imagen, width=Cm(14.2))
    caption = doc.add_paragraph(pie)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.name = "Open Sans Light"; run.font.size = Pt(9); run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(AZUL)


def crear_informe_docx(
    proyecto: dict[str, Any],
    cimentacion: Cimentacion,
    estratos: pd.DataFrame,
    nivel_freatico_m: float,
    z_max_m: float,
    resultado: ResultadoCalculo,
    resultado_steinbrenner: ResultadoSteinbrenner,
) -> bytes:
    """Crea una memoria de cálculo DOCX con estilo geotécnico corporativo."""
    doc = Document()
    _configurar_documento(doc)

    titulo = doc.add_paragraph(style="Title")
    _quitar_bordes(titulo)
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = titulo.add_run("Asiento elástico de zapata rectangular flexible")
    run.font.name = "Open Sans Condensed"; run.font.size = Pt(24); run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(AZUL_OSCURO)
    doc.add_paragraph(
        "Memoria de cálculo del incremento de tensión vertical y de la compresión elástica "
        "de un terreno multicapa bajo una presión rectangular neta uniforme."
    )

    _titulo(doc, "DATOS GENERALES")
    generales = pd.DataFrame(
        [
            ["Proyecto", proyecto.get("proyecto", "")],
            ["Localización", proyecto.get("localizacion", "")],
            ["Autor o revisión", proyecto.get("autor", "")],
            ["Observaciones", proyecto.get("observaciones", "")],
        ], columns=["Campo", "Contenido"]
    )
    _tabla(doc, generales, 2)

    _titulo(doc, "GEOMETRÍA Y MODELO")
    entrada = pd.DataFrame(
        [
            ["Ancho B", cimentacion.b_m, "m"], ["Longitud L", cimentacion.l_m, "m"],
            ["Presión neta", cimentacion.q_neta_kpa, "kPa"], ["Profundidad Df", cimentacion.df_m, "m"],
            ["Nivel freático", nivel_freatico_m, "m desde superficie"],
            ["Profundidad analizada", z_max_m, "m bajo base"],
            ["Parámetros elásticos", "E y ν; M derivado", "-"],
            ["Paso vertical adoptado", resultado.paso_adoptado_m, "m"],
            ["Intervalos adoptados", resultado.intervalos_adoptados, "-"],
            ["Paso vertical de control", resultado.paso_control_m, "m"],
            ["Intervalos de control", resultado.intervalos_control, "-"],
        ], columns=["Parámetro", "Valor", "Unidad"]
    )
    _tabla(doc, entrada)
    doc.add_paragraph(
        "La zapata se considera flexible y transmite una presión neta uniforme. El terreno se "
        "representa mediante capas horizontales elásticas. No se incluyen consolidación, fluencia, "
        "plasticidad ni redistribución de presiones por rigidez de la cimentación."
    )

    doc.add_page_break()
    _titulo(doc, "PARÁMETROS DEL TERRENO")
    _tabla(doc, estratos)

    _titulo(doc, "RESULTADOS POR INTEGRACIÓN VERTICAL")
    asiento_centro = float(resultado.resumen_puntos.loc[resultado.resumen_puntos["Punto"] == "Centro", "asiento_base_mm"].iloc[0])
    doc.add_paragraph(
        f"El asiento elástico calculado bajo el centro es {asiento_centro:.2f} mm para el límite "
        f"inferior situado a {z_max_m:.2f} m bajo la base. El error relativo al duplicar la "
        f"discretización, reduciendo el paso de {resultado.paso_adoptado_m:.3f} m a "
        f"{resultado.paso_control_m:.3f} m, es {resultado.error_discretizacion_pct:.3f} %."
    )
    _tabla(doc, resultado.resumen_puntos)
    _titulo(doc, "APORTE POR ESTRATO", nivel=2)
    _tabla(doc, resultado.resumen_estratos)
    _insertar_figura(doc, _figura_tensiones(resultado), "Figura 1 Perfiles de tensión bajo el centro")
    _insertar_figura(doc, _figura_asientos(resultado), "Figura 2 Desplazamiento vertical en puntos de control")
    _insertar_figura(doc, _figura_bulbo(resultado, cimentacion), "Figura 3 Bulbo de incremento de tensión en la dirección B")

    doc.add_page_break()
    _titulo(doc, "RESULTADOS MEDIANTE STEINBRENNER")
    asiento_stein = float(
        resultado_steinbrenner.resumen_puntos.loc[
            resultado_steinbrenner.resumen_puntos["Punto"] == "Centro", "asiento_base_mm"
        ].iloc[0]
    )
    doc.add_paragraph(
        f"El asiento elástico de Steinbrenner bajo el centro es {asiento_stein:.2f} mm para el "
        f"límite inferior situado a {z_max_m:.2f} m bajo la base. Este cálculo utiliza directamente "
        "E y ν y no depende del paso vertical adoptado para la integración."
    )
    _tabla(doc, resultado_steinbrenner.resumen_puntos)
    _titulo(doc, "APORTE STEINBRENNER POR ESTRATO", nivel=2)
    _tabla(doc, resultado_steinbrenner.resumen_estratos)
    _insertar_figura(
        doc, _figura_steinbrenner(resultado_steinbrenner),
        "Figura 4 Desplazamiento mediante Steinbrenner en puntos de control",
    )

    doc.add_page_break()
    _titulo(doc, "FORMULACIÓN Y LIMITACIONES")
    doc.add_paragraph(
        "El incremento de tensión se obtiene integrando la solución de Boussinesq sobre el área "
        "rectangular. El asiento se calcula como la integral de Δσz/M respecto a la profundidad. "
        "La integral se aproxima mediante una suma evaluada en el punto medio de cada intervalo, "
        "insertando exactamente las interfaces. El módulo confinado se deriva como "
        "M = E(1−ν)/[(1+ν)(1−2ν)]."
    )
    doc.add_paragraph(
        "El cálculo de Steinbrenner se realiza independientemente con los factores I1 e I2, E y ν. "
        "La contribución de cada capa se obtiene como diferencia entre la compresión acumulada en "
        "sus límites superior e inferior. Se aplica superposición rectangular desde el plano de apoyo "
        "con la presión neta, sin factor adicional de empotramiento. Los dos resultados no se combinan "
        "ni se comparan."
    )
    doc.add_paragraph(
        "La distribución de tensiones corresponde a un semiespacio homogéneo; la estratificación "
        "se incorpora en la relación tensión deformación. Los contrastes fuertes de rigidez, la "
        "anisotropía y la respuesta casi incompresible requieren una solución multicapa o numérica."
    )
    doc.add_paragraph(
        f"El último 10 % de la profundidad analizada aporta {resultado.contribucion_tramo_final_pct:.2f} % "
        "del asiento integrado. Este indicador debe revisarse para justificar el límite inferior."
    )

    salida = BytesIO()
    doc.save(salida)
    return salida.getvalue()
