from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from calculos.asientos import ResultadoCalculo
from modelos.datos import Cimentacion


AZUL = "2164AF"
AZUL_OSCURO = "12304A"
AZUL_SUAVE = "DCE8F5"
BLANCO = "FFFFFF"
GRIS = "E8ECF1"


def _formatear_hoja(ws, congelar: str = "A2") -> None:
    ws.freeze_panes = congelar
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for celda in ws[1]:
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.font = Font(name="Arial", size=10, bold=True, color=BLANCO)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borde = Side(style="thin", color="C7CED8")
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name="Arial", size=9, color=AZUL_OSCURO)
            celda.border = Border(bottom=borde)
            if isinstance(celda.value, (int, float)):
                celda.number_format = "0.000"
    for idx, columna in enumerate(ws.columns, start=1):
        longitud = max(len(str(c.value)) if c.value is not None else 0 for c in columna)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longitud + 2, 11), 28)


def crear_excel(
    cimentacion: Cimentacion,
    estratos: pd.DataFrame,
    modo_modulo: str,
    nivel_freatico_m: float,
    z_max_m: float,
    resultado: ResultadoCalculo,
) -> bytes:
    """Genera un libro auditable con entradas y resultados numéricos."""
    salida = BytesIO()
    resumen_caso = pd.DataFrame(
        {
            "Parámetro": ["B", "L", "q neta", "Df", "Nivel freático", "Profundidad analizada", "Modo de rigidez", "Error discretización", "Aporte último 10 %"],
            "Valor": [cimentacion.b_m, cimentacion.l_m, cimentacion.q_neta_kpa, cimentacion.df_m, nivel_freatico_m, z_max_m, modo_modulo, resultado.error_discretizacion_pct, resultado.contribucion_tramo_final_pct],
            "Unidad": ["m", "m", "kPa", "m", "m", "m", "-", "%", "%"],
        }
    )
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen_caso.to_excel(writer, sheet_name="Resumen", index=False, startrow=0)
        resultado.resumen_puntos.to_excel(writer, sheet_name="Resumen", index=False, startrow=len(resumen_caso) + 3)
        estratos.to_excel(writer, sheet_name="Estratos", index=False)
        resultado.resumen_estratos.to_excel(writer, sheet_name="Aporte por estrato", index=False)
        resultado.perfiles.to_excel(writer, sheet_name="Perfiles", index=False)

        for nombre, malla in [("Bulbo eje B", resultado.malla_corta), ("Bulbo eje L", resultado.malla_larga)]:
            datos = pd.DataFrame(malla["delta_sigma_kPa"], index=malla["z_m"], columns=malla["x_m"])
            datos.index.name = "z_bajo_base_m"
            datos.to_excel(writer, sheet_name=nombre)

        for ws in writer.book.worksheets:
            _formatear_hoja(ws)
        ws = writer.book["Resumen"]
        segunda_cabecera = len(resumen_caso) + 4
        for celda in ws[segunda_cabecera]:
            if celda.value is not None:
                celda.fill = PatternFill("solid", fgColor=AZUL)
                celda.font = Font(name="Arial", size=10, bold=True, color=BLANCO)
        ws.merge_cells(start_row=len(resumen_caso) + 2, start_column=1, end_row=len(resumen_caso) + 2, end_column=5)
        titulo = ws.cell(len(resumen_caso) + 2, 1, "Asientos en puntos de control")
        titulo.fill = PatternFill("solid", fgColor=AZUL_SUAVE)
        titulo.font = Font(name="Arial", bold=True, color=AZUL_OSCURO)
    return salida.getvalue()

