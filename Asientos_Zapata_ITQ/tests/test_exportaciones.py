from docx import Document
from openpyxl import load_workbook
import pandas as pd

from calculos.asientos import calcular_caso
from calculos.steinbrenner import calcular_steinbrenner
from modelos.datos import Cimentacion
from servicios.exportacion_docx import crear_informe_docx
from servicios.exportacion_excel import crear_excel


def _caso():
    capas = pd.DataFrame(
        [["Suelo", 0.0, 15.0, 18.0, 20.0, 30.0, 0.30]],
        columns=["Estrato", "z_sup_m", "z_inf_m", "gamma_ap_kN_m3", "gamma_sat_kN_m3", "E_MPa", "nu"],
    )
    cimiento = Cimentacion(2.0, 3.0, 120.0, 1.0)
    resultado = calcular_caso(cimiento, capas, 2.0, 8.0, 0.10)
    resultado_stein = calcular_steinbrenner(cimiento, capas, 8.0)
    capas["M_calculo_MPa"] = 40.38
    return cimiento, capas, resultado, resultado_stein


def test_excel_se_abre_y_contiene_hojas(tmp_path):
    cimiento, capas, resultado, resultado_stein = _caso()
    datos = crear_excel(cimiento, capas, 2.0, 8.0, resultado, resultado_stein)
    ruta = tmp_path / "resultado.xlsx"; ruta.write_bytes(datos)
    wb = load_workbook(ruta, read_only=True, data_only=False)
    assert {
        "Resumen", "Estratos", "Integración - Puntos", "Integración - Estratos",
        "Steinbrenner - Puntos", "Steinbrenner - Estratos", "Bulbo eje B", "Bulbo eje L",
    } <= set(wb.sheetnames)


def test_docx_se_abre_y_contiene_resultado(tmp_path):
    cimiento, capas, resultado, resultado_stein = _caso()
    datos = crear_informe_docx(
        {"proyecto": "Prueba"}, cimiento, capas, 2.0, 8.0,
        resultado, resultado_stein,
    )
    ruta = tmp_path / "resultado.docx"; ruta.write_bytes(datos)
    doc = Document(ruta)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "Asiento elástico" in texto
    assert "RESULTADOS POR INTEGRACIÓN VERTICAL" in texto
    assert "RESULTADOS MEDIANTE STEINBRENNER" in texto
