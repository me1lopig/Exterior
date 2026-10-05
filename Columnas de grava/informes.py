import pandas as pd
import io
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import xlsxwriter

def generar_excel_memoria(df_geo, df_int1, df_int2, df_params, inputs_dict):
    """
    Genera un archivo Excel multipestaña en memoria (BytesIO) listo para descarga.
    """
    output = io.BytesIO()
    
    # Usamos xlsxwriter como motor para permitir formateos si fuera necesario
    with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
        # Pestaña 1: Inputs
        df_inputs = pd.DataFrame(list(inputs_dict.items()), columns=['Parámetro', 'Valor'])
        df_inputs.to_excel(writer, sheet_name='1_Inputs', index=False)
        
        # Pestañas de cálculo
        df_geo.to_excel(writer, sheet_name='2_Geometria', index=False)
        
        # Unimos los intermedios por simplicidad de lectura
        df_intermedios = pd.merge(df_int1, df_int2, on='Diámetro Columa de grava (m)')
        df_intermedios.to_excel(writer, sheet_name='3_Priebe_Intermedios', index=False)
        
        df_params.to_excel(writer, sheet_name='4_Parametros_Equivalentes', index=False)
        
    return output.getvalue()

def generar_word_memoria(df_params, figs, inputs_dict):
    """
    Genera la Memoria de Cálculo en Word. 
    figs: Lista de figuras Plotly ya generadas.
    """
    doc = Document()
    
    # Estilo del título
    titulo = doc.add_heading('MEMORIA DE CÁLCULO: COLUMNAS DE GRAVA', level=1)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('1. Objeto y Bases de Cálculo', level=2)
    doc.add_paragraph(
        "El presente anexo justifica la homogeneización paramétrica de un macizo de suelo "
        "tratado mediante inclusiones granulares (columnas de grava). La estimación de la mejora "
        "tensional y el factor de reducción de asientos se evalúa bajo los postulados del Método de Priebe (1995)."
    )
    
    doc.add_heading('2. Parámetros de Diseño', level=2)
    for k, v in inputs_dict.items():
        doc.add_paragraph(f"{k}: {v}", style='List Bullet')
        
    doc.add_heading('3. Evolución de Parámetros Equivalentes', level=2)
    doc.add_paragraph("A continuación, se presentan las gráficas de sensibilidad analizando la variación tensional en función del diámetro de la inclusión:")
    
    # Insertar gráficos (requiere kaleido instalado para fig.write_image)
    for i, fig in enumerate(figs):
        img_bytes = fig.to_image(format="png", width=700, height=400, scale=2)
        image_stream = io.BytesIO(img_bytes)
        doc.add_picture(image_stream, width=Inches(6.0))
        ultimo_parrafo = doc.paragraphs[-1]
        ultimo_parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    # Guardar en buffer
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
