"""Generacion de un informe tecnico DOCX con el estilo corporativo disponible."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from servicios.catalogo import CORRELACIONES, FUENTES
from servicios.version import APP_VERSION


AZUL = "2164AF"
ESTILOS_PATH = Path(__file__).resolve().parents[1] / "agent_files" / "estilos_informes_geotecnicos.json"


def _borde_inferior(parrafo: object) -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement("w:pBdr")
    borde = OxmlElement("w:bottom")
    borde.set(qn("w:val"), "single")
    borde.set(qn("w:sz"), "8")
    borde.set(qn("w:color"), AZUL)
    bordes.append(borde)
    p_pr.append(bordes)


def _configurar(documento: Document) -> None:
    with ESTILOS_PATH.open(encoding="utf-8") as archivo:
        estilos = json.load(archivo)
    configuracion = estilos["configuracion_documento"]["secciones"][0]
    margenes = configuracion["margenes_mm"]
    azul = estilos["resumen"]["paleta_corporativa_identificada"]["azul_principal"].lstrip("#")
    seccion = documento.sections[0]
    seccion.page_width = Mm(configuracion["ancho_mm"])
    seccion.page_height = Mm(configuracion["alto_mm"])
    seccion.top_margin = Cm(margenes["top"] / 10)
    seccion.bottom_margin = Cm(margenes["bottom"] / 10)
    seccion.left_margin = Cm(margenes["left"] / 10)
    seccion.right_margin = Cm(margenes["right"] / 10)
    normal = documento.styles["Normal"]
    normal.font.name = "Open Sans"
    normal.font.size = Pt(9)
    titulo = documento.styles["Title"]
    titulo.font.name = "Open Sans"
    titulo.font.size = Pt(20)
    titulo.font.bold = True
    titulo.font.color.rgb = RGBColor.from_string("10263F")
    for nombre in ("Heading 1", "Heading 2"):
        estilo = documento.styles[nombre]
        estilo.font.name = "Open Sans"
        estilo.font.color.rgb = RGBColor.from_string(azul)
        estilo.font.bold = True


def _titulo(documento: Document, texto: str, nivel: int = 1) -> None:
    parrafo = documento.add_heading(texto.upper() if nivel == 1 else texto, level=nivel)
    if nivel == 1:
        _borde_inferior(parrafo)


def _tabla(documento: Document, df: pd.DataFrame) -> None:
    if df.empty:
        documento.add_paragraph("Sin datos.")
        return
    def texto(valor: object) -> str:
        if pd.isna(valor):
            return ""
        if isinstance(valor, float):
            return f"{valor:.6g}"
        return str(valor)

    vista = df.map(texto)
    tabla = documento.add_table(rows=1, cols=len(vista.columns))
    tabla.style = "Table Grid"
    cabecera_pr = tabla.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    cabecera_pr.append(repetir)
    for i, columna in enumerate(vista.columns):
        tabla.rows[0].cells[i].text = str(columna)
    for _, fila in vista.iterrows():
        fila_tabla = tabla.add_row()
        fila_pr = fila_tabla._tr.get_or_add_trPr()
        no_dividir = OxmlElement("w:cantSplit")
        fila_pr.append(no_dividir)
        celdas = fila_tabla.cells
        for i, valor in enumerate(fila):
            celdas[i].text = valor
    for indice_fila, fila in enumerate(tabla.rows):
        for celda in fila.cells:
            for parrafo in celda.paragraphs:
                parrafo.paragraph_format.space_after = Pt(0)
                for run in parrafo.runs:
                    run.font.name = "Open Sans"
                    run.font.size = Pt(7.5)
                    if indice_fila == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if indice_fila == 0:
                propiedades = celda._tc.get_or_add_tcPr()
                sombreado = OxmlElement("w:shd")
                sombreado.set(qn("w:fill"), AZUL)
                propiedades.append(sombreado)


def generar_informe(
    proyecto: str,
    localizacion: str,
    grupos: list[dict[str, str]],
    resultados: pd.DataFrame,
    entradas: dict[str, pd.DataFrame] | None = None,
) -> bytes:
    documento = Document()
    _configurar(documento)
    titulo = documento.add_paragraph(style="Title")
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    titulo.add_run("Estimacion de parametros geotecnicos")
    documento.add_paragraph(f"Proyecto: {proyecto}")
    documento.add_paragraph(f"Localizacion: {localizacion or 'No indicada'}")
    documento.add_paragraph(f"Version de la aplicacion: {APP_VERSION}")
    documento.add_paragraph(f"Fecha UTC del expediente: {datetime.now(timezone.utc).isoformat(timespec='seconds')}")

    _titulo(documento, "Alcance")
    documento.add_paragraph(
        "Estimacion mediante correlaciones basadas en ensayos de campo y valores "
        "orientativos del CTE DB-SE-C. Los resultados derivados no sustituyen la "
        "seleccion del valor caracteristico por el tecnico responsable."
    )
    _titulo(documento, "Grupos de datos")
    _tabla(documento, pd.DataFrame(grupos))
    _titulo(documento, "Datos de entrada")
    hay_entradas = False
    for tipo, datos in (entradas or {}).items():
        if datos.empty:
            continue
        hay_entradas = True
        _titulo(documento, tipo.upper(), nivel=2)
        _tabla(documento, datos)
    if not hay_entradas:
        documento.add_paragraph("No se han incorporado tablas de entrada al expediente.")
    _titulo(documento, "Resultados")
    columnas = [
        c for c in ["grupo", "punto", "profundidad_m", "parametro", "valor", "minimo", "maximo",
                    "unidad", "estado"] if c in resultados.columns
    ]
    _tabla(documento, resultados[columnas] if columnas else resultados)
    _titulo(documento, "Trazabilidad por resultado", nivel=2)
    columnas_trazabilidad = [
        c for c in ["grupo", "punto", "parametro", "correlacion_id", "categoria", "interpolado",
                    "advertencias"] if c in resultados.columns
    ]
    _tabla(documento, resultados[columnas_trazabilidad] if columnas_trazabilidad else pd.DataFrame())
    _titulo(documento, "Valores adoptados")
    if "adoptado" in resultados and resultados["adoptado"].fillna(False).any():
        adoptados = resultados[resultados["adoptado"].fillna(False)].copy()
        columnas_adoptadas = [
            c for c in ["grupo", "unidad_geotecnica", "parametro", "valor_adoptado", "unidad",
                        "correlacion_id", "estado", "justificacion"] if c in adoptados
        ]
        _tabla(documento, adoptados[columnas_adoptadas])
    else:
        documento.add_paragraph("No se han seleccionado valores adoptados.")
    _titulo(documento, "Metodos y fuentes")
    usados = sorted(set(resultados.get("correlacion_id", pd.Series(dtype=str)).dropna()))
    trazabilidad = []
    for correlacion_id in usados:
        correlacion = CORRELACIONES.get(str(correlacion_id))
        if correlacion is None:
            continue
        fuente = FUENTES[correlacion.fuente_id]
        trazabilidad.append({
            "ID": correlacion.id,
            "Version": correlacion.version_metodo,
            "Formula": correlacion.formula,
            "Dominio": correlacion.dominio,
            "Fuente": fuente.referencia,
            "URL": fuente.url,
        })
    _tabla(documento, pd.DataFrame(trazabilidad))
    _titulo(documento, "Limitaciones")
    documento.add_paragraph(
        "Las correlaciones empiricas conservan el dominio y las limitaciones de su fuente. "
        "No se realizan extrapolaciones fuera de los intervalos implementados. La interpolacion "
        "lineal de D.23 es un criterio adoptado por GeoCorr y no una prescripcion del CTE. "
        "Los resultados no son automaticamente valores caracteristicos, de calculo o admisibles."
    )
    salida = BytesIO()
    documento.save(salida)
    return salida.getvalue()
