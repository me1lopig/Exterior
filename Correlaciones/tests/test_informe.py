import unittest
from io import BytesIO

import pandas as pd
from docx import Document

from servicios.informes_docx import generar_informe


class TestInforme(unittest.TestCase):
    def test_docx_se_genera_y_reabre(self) -> None:
        resultados = pd.DataFrame([{
            "grupo": "G01",
            "punto": "S-1",
            "profundidad_m": 3.0,
            "correlacion_id": "RES-PHI-SPT-CTE2019-T41",
            "parametro": "phi_eff",
            "metodo": "CTE 4.1",
            "valor": 34.0,
            "minimo": None,
            "maximo": None,
            "unidad": "grados",
            "advertencias": "",
            "estado": "Valido",
            "interpolado": True,
            "adoptado": True,
            "valor_adoptado": 33.5,
            "justificacion": "Valor conservador dentro del intervalo.",
        }])
        entradas = {"spt": pd.DataFrame([{
            "grupo": "G01", "punto": "S-1", "tipo_n": "NSPT compatible CTE", "n": 22,
        }])}
        contenido = generar_informe(
            "Prueba", "Madrid", [{"id": "G01", "nombre": "S-1", "ensayo": "SPT"}],
            resultados, entradas=entradas,
        )
        documento = Document(BytesIO(contenido))
        self.assertGreater(len(contenido), 10_000)
        self.assertGreaterEqual(len(documento.tables), 2)
        self.assertIn("Estimacion de parametros", documento.paragraphs[0].text)
        self.assertAlmostEqual(documento.sections[0].page_width.cm, 21.0009, places=2)
        self.assertAlmostEqual(documento.sections[0].page_height.cm, 29.7004, places=2)
        texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)
        self.assertIn("DATOS DE ENTRADA", texto)
        self.assertIn("VALORES ADOPTADOS", texto)
        self.assertIn("METODOS Y FUENTES", texto)


if __name__ == "__main__":
    unittest.main()
