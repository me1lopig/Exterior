from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from pathlib import PurePath
from typing import BinaryIO, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook

from modelos.datos import DatosProyecto, EstadoNF


class ErrorDatos(ValueError):
    """Error comprensible para el usuario durante la importacion."""


def _normalizar(text: object) -> str:
    value = "" if text is None else str(text).strip().lower()
    value = "".join(
        char for char in unicodedata.normalize("NFKD", value) if not unicodedata.combining(char)
    )
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def expandir_archivos(uploaded: Iterable[object]) -> dict[str, bytes]:
    """Devuelve XLSX individuales, incluyendo los contenidos en ZIP seguros."""
    result: dict[str, bytes] = {}
    for item in uploaded:
        name = PurePath(str(getattr(item, "name", "archivo"))).name
        content = item.getvalue() if hasattr(item, "getvalue") else bytes(item)
        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    for member in archive.infolist():
                        safe_name = PurePath(member.filename).name
                        if member.is_dir() or not safe_name.lower().endswith(".xlsx"):
                            continue
                        if member.file_size > 25 * 1024 * 1024:
                            raise ErrorDatos(f"{safe_name}: supera el limite de 25 MB.")
                        if safe_name.lower() in {key.lower() for key in result}:
                            raise ErrorDatos(f"Nombre XLSX duplicado: {safe_name}.")
                        result[safe_name] = archive.read(member)
            except zipfile.BadZipFile as exc:
                raise ErrorDatos(f"{name} no es un ZIP valido.") from exc
        elif name.lower().endswith(".xlsx"):
            if name.lower() in {key.lower() for key in result}:
                raise ErrorDatos(f"Nombre XLSX duplicado: {name}.")
            result[name] = content
    return result


def _rows(content: bytes) -> tuple[list[str], list[tuple[object, ...]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ErrorDatos("No se pudo abrir uno de los libros XLSX.") from exc
    sheet = workbook.worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not values:
        return [], []
    headers = ["" if value is None else str(value).strip() for value in values[0]]
    return headers, values[1:]


def _find(files: dict[str, bytes], expected: str, *, required: bool = True) -> bytes | None:
    for name, content in files.items():
        if name.casefold() == expected.casefold():
            return content
    if required:
        raise ErrorDatos(f"Falta el archivo requerido {expected}.")
    return None


def _cell(row: tuple[object, ...], index: int) -> object:
    return row[index] if index < len(row) else None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _profile_schema(headers: list[str]) -> bool:
    normalized = {_normalizar(header) for header in headers}
    return {"estacion", "distancia", "cota"}.issubset(normalized)


def _profile_columns(headers: list[str]) -> tuple[int, int, int]:
    normalized = [_normalizar(header) for header in headers]
    return normalized.index("estacion"), normalized.index("distancia"), normalized.index("cota")


def _dataframe(records: list[dict[str, object]], columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame.from_records(records, columns=columns)


def cargar_proyecto(files: dict[str, bytes]) -> DatosProyecto:
    bkfl_headers, bkfl_rows = _rows(_find(files, "BKFL.XLSX"))
    dprg_headers, dprg_rows = _rows(_find(files, "DPRG.XLSX"))
    lab_headers, lab_rows = _rows(_find(files, "ListadoLab.XLSX"))
    ispt_headers, ispt_rows = _rows(_find(files, "ISPT.XLSX"))
    ugeo_content = _find(files, "UGEO.XLSX", required=False)
    ugeo_headers, ugeo_rows = _rows(ugeo_content) if ugeo_content else ([], [])

    minimum_columns = {
        "BKFL.XLSX": (bkfl_headers, 7),
        "DPRG.XLSX": (dprg_headers, 8),
        "ListadoLab.XLSX": (lab_headers, 40),
        "ISPT.XLSX": (ispt_headers, 7),
    }
    for name, (headers, minimum) in minimum_columns.items():
        if len(headers) < minimum:
            raise ErrorDatos(f"{name}: se esperaban al menos {minimum} columnas y hay {len(headers)}.")

    litologia = _dataframe(
        [
            {
                "sondeo": _text(_cell(row, 0)),
                "profundidad_inicial_m": _cell(row, 2),
                "profundidad_final_m": _cell(row, 3),
                "descripcion_corta": _text(_cell(row, 5)),
                "descripcion": _text(_cell(row, 6)),
            }
            for row in bkfl_rows
            if _text(_cell(row, 0))
        ],
        ["sondeo", "profundidad_inicial_m", "profundidad_final_m", "descripcion_corta", "descripcion"],
    )
    dpsh = _dataframe(
        [
            {
                "ensayo": _text(_cell(row, 0)),
                "profundidad_m": _cell(row, 6),
                "golpes": _cell(row, 7),
            }
            for row in dprg_rows
            if _text(_cell(row, 0))
        ],
        ["ensayo", "profundidad_m", "golpes"],
    )
    ensayos = _dataframe(
        [
            {
                "sondeo": _text(_cell(row, 1)),
                "tipo": _text(_cell(row, 30)) or _text(_cell(row, 2)),
                "profundidad_m": _cell(row, 3),
                "n_spt": _text(_cell(row, 8)),
                "n_mi": _text(_cell(row, 9)),
                "uscs": _text(_cell(row, 21)),
                "rcs_kpa": _text(_cell(row, 36)),
                "phi_deg": _text(_cell(row, 38)),
                "cohesion_kpa": _text(_cell(row, 39)),
            }
            for row in lab_rows
            if _text(_cell(row, 1))
        ],
        ["sondeo", "tipo", "profundidad_m", "n_spt", "n_mi", "uscs", "rcs_kpa", "phi_deg", "cohesion_kpa"],
    )
    ugeo = _dataframe(
        [
            {
                "sondeo": _text(_cell(row, 0)),
                "profundidad_inicial_m": _cell(row, 2),
                "profundidad_final_m": _cell(row, 3),
                "codigo": _text(_cell(row, 4)),
                "descripcion": _text(_cell(row, 5)),
            }
            for row in ugeo_rows
            if _text(_cell(row, 0))
        ],
        ["sondeo", "profundidad_inicial_m", "profundidad_final_m", "codigo", "descripcion"],
    )

    nf_records: list[dict[str, object]] = []
    seen: set[str] = set()
    for row in ispt_rows:
        name = _text(_cell(row, 0))
        if not name or name in seen:
            continue
        seen.add(name)
        value = _cell(row, 6)
        if value in (None, ""):
            state, depth = EstadoNF.NO_MEDIDO.value, None
        else:
            try:
                depth = float(value)
            except (TypeError, ValueError):
                state, depth = EstadoNF.NO_MEDIDO.value, None
            else:
                state = EstadoNF.REVISAR.value if depth == 0 else EstadoNF.DETECTADO.value
        nf_records.append({"sondeo": name, "estado": state, "profundidad_m": depth})
    nivel_freatico = _dataframe(nf_records, ["sondeo", "estado", "profundidad_m"])

    known_boreholes = set(litologia["sondeo"].astype(str))
    known_dpsh = set(dpsh["ensayo"].astype(str))
    profile_records: list[dict[str, object]] = []
    ignored: list[str] = []
    fixed_names = {"bkfl.xlsx", "dprg.xlsx", "listadolab.xlsx", "ispt.xlsx", "ugeo.xlsx"}
    for name, content in files.items():
        if name.casefold() in fixed_names:
            continue
        headers, rows = _rows(content)
        if not _profile_schema(headers):
            ignored.append(name)
            continue
        station_col, distance_col, elevation_col = _profile_columns(headers)
        profile_name = PurePath(name).stem
        for order, row in enumerate(rows):
            station = _text(_cell(row, station_col))
            if not station:
                continue
            if station in known_boreholes:
                station_type = "Sondeo"
            elif station in known_dpsh:
                station_type = "DPSH"
            else:
                station_type = "Punto de terreno"
            profile_records.append(
                {
                    "perfil": profile_name,
                    "orden": order,
                    "estacion": station,
                    "tipo": station_type,
                    "distancia_m": _cell(row, distance_col),
                    "cota_m": _cell(row, elevation_col),
                }
            )
    perfiles = _dataframe(
        profile_records, ["perfil", "orden", "estacion", "tipo", "distancia_m", "cota_m"]
    )
    if perfiles.empty:
        raise ErrorDatos("No se encontro ningun XLSX con columnas ESTACION, Distancia y Cota.")

    return DatosProyecto(
        perfiles=perfiles,
        litologia=litologia,
        dpsh=dpsh,
        ensayos=ensayos,
        ugeo=ugeo,
        nivel_freatico=nivel_freatico,
        archivos_ignorados=ignored,
    )


def exportar_excel_corregido(datos: DatosProyecto, original: DatosProyecto | None = None) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    export_data = datos.copy()
    if not export_data.nivel_freatico.empty:
        not_detected = export_data.nivel_freatico["estado"] != EstadoNF.DETECTADO.value
        export_data.nivel_freatico.loc[not_detected, "profundidad_m"] = None
    for name, frame in export_data.tablas().items():
        sheet = workbook.create_sheet(name[:31])
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    if original is not None:
        changes = workbook.create_sheet("Cambios")
        changes.append(["Tabla", "Fila", "Campo", "Valor original", "Valor corregido"])
        original_tables = original.tablas()
        for table_name, current in export_data.tablas().items():
            before = original_tables[table_name]
            columns = list(dict.fromkeys([*before.columns, *current.columns]))
            for row_index in range(max(len(before), len(current))):
                for column in columns:
                    old = before.iloc[row_index].get(column) if row_index < len(before) else None
                    new = current.iloc[row_index].get(column) if row_index < len(current) else None
                    both_empty = pd.isna(old) and pd.isna(new)
                    if both_empty or (not pd.isna(old) and not pd.isna(new) and old == new):
                        continue
                    changes.append(
                        [
                            table_name,
                            row_index + 2,
                            column,
                            None if pd.isna(old) else old,
                            None if pd.isna(new) else new,
                        ]
                    )
        changes.freeze_panes = "A2"
        changes.auto_filter.ref = changes.dimensions
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
