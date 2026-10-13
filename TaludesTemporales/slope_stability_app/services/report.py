"""Informe DOCX autocontenido de un analisis de estabilidad."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from slope_stability_app.calculations.geometry import ground_elevation
from slope_stability_app.calculations.search import SearchResult
from slope_stability_app.models import AnalysisConfig, ProjectData, SlopeGeometry, SoilProfile


NAVY = "263746"
PALE = "EAF0F4"
GRID = "D9D9D9"
AMBER = "D97706"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRID)
        borders.append(node)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = heading
        _shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cells[col_index].text = str(value)
            cells[col_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                _shade(cells[col_index], "F6F8FA")
            for paragraph in cells[col_index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        if widths:
            for cell, width in zip(cells, widths):
                cell.width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _section_image(geometry: SlopeGeometry, profile: SoilProfile, result: SearchResult) -> BytesIO:
    critical = result.bishop
    circle = critical.evaluation.circle
    x = np.linspace(geometry.x_min, geometry.x_max, 500)
    ground = ground_elevation(x, geometry)
    fig, ax = plt.subplots(figsize=(8.1, 4.0), dpi=180)
    bottom = -0.9 * geometry.height
    if profile.has_two_soils:
        interface = profile.interface_elevation
        ax.fill_between(x, bottom, np.minimum(ground, interface), color="#CFC1AA", alpha=0.72)
        upper_mask = ground >= interface
        ax.fill_between(
            x[upper_mask],
            interface,
            ground[upper_mask],
            color="#E6DDD0",
            alpha=0.90,
        )
    else:
        ax.fill_between(x, bottom, ground, color="#E6DDD0", alpha=0.85)
    ax.plot(x, ground, color="#17202A", lw=1.8, label="Terreno")
    water_mask = ground >= geometry.water_elevation - 1e-9
    ax.plot(
        x[water_mask],
        np.full(np.count_nonzero(water_mask), geometry.water_elevation),
        color="#2878B5",
        lw=1.6,
        ls="--",
        label="Nivel freatico",
    )
    if profile.has_two_soils:
        contact_mask = ground >= profile.interface_elevation
        ax.plot(
            x[contact_mask],
            np.full(np.count_nonzero(contact_mask), profile.interface_elevation),
            color="#8B5E3C",
            lw=1.4,
            ls=":",
            label="Contacto",
        )
    xs = np.linspace(circle.entry_x, circle.exit_x, 350)
    ax.plot(xs, circle.base_elevation(xs), color="#D97706", lw=2.2, label=f"Bishop FS={critical.factor_of_safety:.3f}")
    for item in critical.evaluation.slices:
        ax.plot([item.x_mid, item.x_mid], [item.base_elevation, item.top_elevation], color="#65717B", lw=0.35, alpha=0.6)
    ax.scatter([circle.center_x], [circle.center_y], marker="+", color="#D97706", s=50)
    ax.set_xlabel("Distancia x (m)")
    ax.set_ylabel("Cota z (m)")
    ax.set_aspect("equal", adjustable="datalim")
    ax.grid(color="#E7E9ED", lw=0.6)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, 1.17), ncol=4, frameon=False, fontsize=8)
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _sensitivity_image(data: pd.DataFrame) -> BytesIO:
    fig, ax = plt.subplots(figsize=(8.1, 3.8), dpi=180)
    angles = data["Pendiente (deg)"].nunique()
    waters = data["Nivel freatico (m)"].nunique()
    if angles > 1 and waters > 1:
        pivot = data.pivot(index="Nivel freatico (m)", columns="Pendiente (deg)", values="FS Bishop")
        image = ax.imshow(pivot.values, origin="upper", aspect="auto", cmap="RdYlGn")
        ax.set_xticks(range(len(pivot.columns)), [f"{v:.1f}" for v in pivot.columns])
        ax.set_yticks(range(len(pivot.index)), [f"{v:.2f}" for v in pivot.index])
        ax.set_xlabel("Pendiente (deg)")
        ax.set_ylabel("Cota NF (m)")
        fig.colorbar(image, ax=ax, label="FS Bishop")
    else:
        variable = "Pendiente (deg)" if angles > 1 else "Nivel freatico (m)"
        ax.plot(data[variable], data["FS Fellenius"], "o-", label="Fellenius")
        ax.plot(data[variable], data["FS Bishop"], "o-", color="#D97706", label="Bishop simplificado")
        ax.set_xlabel(variable)
        ax.set_ylabel("Factor de seguridad")
        ax.legend(frameon=False)
    ax.grid(color="#E7E9ED", lw=0.6)
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Title", 23), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Analisis preliminar de estabilidad de talud temporal | Herramienta GeoTech").font.size = Pt(8)


def build_report(
    project: ProjectData,
    geometry: SlopeGeometry,
    profile: SoilProfile,
    config: AnalysisConfig,
    result: SearchResult,
    sensitivity: pd.DataFrame | None = None,
) -> bytes:
    """Genera un DOCX en memoria listo para descarga."""
    document = Document()
    _configure_document(document)
    title = document.add_paragraph(style="Title")
    title.add_run("Estabilidad de talud temporal")
    subtitle = document.add_paragraph()
    subtitle.add_run(project.name or "Analisis sin titulo").bold = True
    subtitle.add_run(f"\n{project.location}" if project.location else "")
    subtitle.add_run(f"\nFecha: {datetime.now():%d/%m/%Y}")
    if project.author:
        subtitle.add_run(f"  |  Autor: {project.author}")
    document.add_paragraph(
        "Informe de calculo por equilibrio limite para superficies circulares. Los resultados son una ayuda al proyecto y no sustituyen la definicion del modelo geotecnico, las combinaciones de acciones ni las comprobaciones exigibles en el marco normativo aplicable."
    )

    document.add_heading("1. Resumen de resultados", level=1)
    governing = min(result.fellenius.factor_of_safety, result.bishop.factor_of_safety)
    status = "Cumple el objetivo introducido" if governing >= config.target_factor else "No alcanza el objetivo introducido"
    _add_table(
        document,
        ["Metodo", "FS minimo", "Centro x (m)", "Centro z (m)", "Radio (m)", "Evaluacion"],
        [
            [
                critical.method,
                f"{critical.factor_of_safety:.3f}",
                f"{critical.evaluation.circle.center_x:.2f}",
                f"{critical.evaluation.circle.center_y:.2f}",
                f"{critical.evaluation.circle.radius:.2f}",
                "FS >= objetivo" if critical.factor_of_safety >= config.target_factor else "FS < objetivo",
            ]
            for critical in (result.fellenius, result.bishop)
        ],
    )
    document.add_paragraph(
        f"Resultado gobernante: FS = {governing:.3f}. Factor objetivo definido por el usuario: {config.target_factor:.2f}. {status}. Se evaluaron {result.evaluated_surfaces} circulos, de los cuales {result.valid_surfaces} fueron geometricamente y numericamente validos."
    )
    document.add_picture(_section_image(geometry, profile, result), width=Inches(6.75))
    caption = document.add_paragraph("Figura 1. Geometria y superficie critica obtenida con Bishop simplificado.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8.5)

    document.add_page_break()
    document.add_heading("2. Datos de entrada", level=1)
    _add_table(
        document,
        ["Parametro", "Valor", "Unidad / criterio"],
        [
            ["Altura H", f"{geometry.height:.2f}", "m"],
            ["Angulo del talud", f"{geometry.angle_deg:.2f}", "grados respecto a horizontal"],
            ["Desarrollo horizontal", f"{geometry.run:.2f}", "m"],
            ["Cota del nivel freatico", f"{geometry.water_elevation:.2f}", "m sobre el pie"],
            ["Numero de dovelas", str(config.slices), "-"],
            ["Busqueda", config.search_level, "malla determinista"],
            ["Peso unitario del agua", f"{config.gamma_water:.2f}", "kN/m3"],
            ["Factor objetivo", f"{config.target_factor:.2f}", "definido por el usuario"],
        ],
    )
    soil_rows = []
    for position, soil in (("Superior", profile.upper), ("Inferior", profile.lower)):
        if soil is None:
            continue
        soil_rows.append([position, soil.name, f"{soil.gamma_dry:.2f}", f"{soil.gamma_sat:.2f}", f"{soil.cohesion:.2f}", f"{soil.friction_deg:.2f}"])
    _add_table(document, ["Estrato", "Suelo", "gamma seco", "gamma sat.", "c' (kPa)", "phi' (deg)"], soil_rows)
    if profile.has_two_soils:
        document.add_paragraph(
            f"El estrato superior tiene un espesor vertical de {profile.upper_thickness(geometry.height):.2f} m medido desde la coronacion. "
            f"El contacto horizontal se situa en z = {profile.interface_elevation:.2f} m sobre el pie; entre el contacto y el pie hay "
            f"{profile.lower_thickness_within_slope(geometry.height):.2f} m del estrato inferior, que se considera prolongado bajo z=0."
        )

    document.add_heading("3. Metodo e hipotesis", level=1)
    assumptions = [
        "Analisis bidimensional por metro de ancho, deformacion plana y mecanismo de rotura circular.",
        "Resistencia drenada Mohr-Coulomb con c' y phi' constantes dentro de cada estrato.",
        "Fellenius desprecia las fuerzas entre dovelas; Bishop simplificado satisface el equilibrio de momentos e itera el factor de seguridad.",
        "La presion intersticial es hidrostatica bajo el nivel freatico horizontal y nula sobre el mismo. El nivel se recorta contra el terreno, sin modelar agua exterior ni fuerzas de filtracion adicionales.",
        "Los pesos se integran por estrato y condicion seca/saturada. La resistencia de la base de cada dovela corresponde al suelo atravesado por su punto medio.",
        "La superficie critica es el minimo dentro de la familia discreta explorada; una busqueda mas densa o una geometria real pueden producir un minimo inferior.",
    ]
    for item in assumptions:
        document.add_paragraph(item, style="List Bullet")

    if sensitivity is not None and not sensitivity.empty:
        document.add_section(WD_SECTION_START.NEW_PAGE)
        document.add_heading("4. Sensibilidad parametrica", level=1)
        minimum_row = sensitivity.loc[sensitivity["FS Bishop"].idxmin()]
        maximum_row = sensitivity.loc[sensitivity["FS Bishop"].idxmax()]
        document.add_paragraph(
            f"Se analizaron {len(sensitivity)} combinaciones. Bishop varia entre {minimum_row['FS Bishop']:.3f} y {maximum_row['FS Bishop']:.3f}. El minimo corresponde a pendiente {minimum_row['Pendiente (deg)']:.1f} grados y nivel freatico z={minimum_row['Nivel freatico (m)']:.2f} m."
        )
        document.add_picture(_sensitivity_image(sensitivity), width=Inches(6.7))
        display = sensitivity.copy()
        if len(display) > 15:
            display = display.sort_values("FS Bishop").head(15)
            document.add_paragraph("La tabla muestra las 15 combinaciones mas desfavorables; el CSV descargable conserva todos los casos.")
        rows = [
            [f"{row['Pendiente (deg)']:.2f}", f"{row['Nivel freatico (m)']:.2f}", f"{row['FS Fellenius']:.3f}", f"{row['FS Bishop']:.3f}"]
            for _, row in display.iterrows()
        ]
        _add_table(document, ["Pendiente (deg)", "NF (m)", "FS Fellenius", "FS Bishop"], rows)

    document.add_heading("5. Limitaciones y recomendaciones", level=1)
    document.add_paragraph(
        "Antes de adoptar el resultado deben contrastarse la estratigrafia, el regimen piezometrico, la representatividad de c' y phi', la posible respuesta no drenada, las sobrecargas, las excavaciones por fases y los mecanismos no circulares. El factor objetivo debe fijarse expresamente segun normativa, duracion de la situacion temporal, fiabilidad de la investigacion y consecuencias de fallo."
    )
    if project.notes:
        document.add_heading("6. Observaciones del proyecto", level=1)
        document.add_paragraph(project.notes)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
