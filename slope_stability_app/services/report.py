"""Informe DOCX autocontenido de un analisis de estabilidad."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from slope_stability_app.calculations.geometry import ground_elevation
from slope_stability_app.calculations.search import SearchResult
from slope_stability_app.models import AnalysisConfig, ProjectData, SlopeGeometry, SoilProfile


GEO_BLUE = "2164AF"
PALE = "F2F6FB"
GRID = "D9D9D9"
AMBER = "D97706"
SOIL_COLORS = ("#E0D2B4", "#B8CABA", "#C7BED3", "#DABEAE", "#B1C6D3", "#D3C9A9")


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_margins(cell, top: int = 80, start: int = 70, bottom: int = 80, end: int = 70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6" if edge in ("top", "bottom") else "3")
        node.set(qn("w:color"), GEO_BLUE if edge in ("top", "bottom") else GRID)
        borders.append(node)


def _append_seq_field(paragraph, sequence: str, number: int) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" SEQ {sequence} \\* ARABIC "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = str(number)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, cached, end))


def _add_caption(document: Document, kind: str, number: int, text: str):
    style = "GEO pie de figura" if kind == "Figura" else "GEO tabla titulo"
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run("Figura N. " if kind == "Figura" else "Tabla ")
    _append_seq_field(paragraph, kind, number)
    paragraph.add_run(".- " if kind == "Figura" else ". ")
    paragraph.add_run(text)
    return paragraph


def _add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    title: str,
    number: int,
    widths: list[float] | None = None,
):
    _add_caption(document, "Tabla", number, title)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = heading
        _shade(cell, PALE)
        _set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.name = "Open Sans"
            run.font.color.rgb = RGBColor.from_string(GEO_BLUE)
            run.font.size = Pt(9)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cells[col_index].text = str(value)
            cells[col_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[col_index])
            if row_index % 2:
                _shade(cells[col_index], PALE)
            for paragraph in cells[col_index].paragraphs:
                paragraph.style = document.styles["GEO tabla contenido"]
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
        if widths:
            for cell, width in zip(cells, widths):
                cell.width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _section_image(geometry: SlopeGeometry, profile: SoilProfile, result: SearchResult) -> BytesIO:
    critical = result.bishop
    circle = critical.evaluation.circle
    x = np.linspace(geometry.x_min, geometry.x_max, 500)
    ground = ground_elevation(x, geometry)
    fig, ax = plt.subplots(figsize=(8.1, 4.0), dpi=180)
    bottom = geometry.model_base_elevation
    cumulative = 0.0
    for index, layer in enumerate(profile.layers):
        layer_top = geometry.height - cumulative
        cumulative += layer.thickness
        layer_bottom = max(geometry.height - cumulative, bottom)
        if layer_top <= bottom + 1e-9:
            break
        layer_mask = ground >= layer_bottom - 1e-9
        ax.fill_between(
            x[layer_mask],
            layer_bottom,
            np.minimum(ground[layer_mask], layer_top),
            color=SOIL_COLORS[index % len(SOIL_COLORS)],
            alpha=0.78,
            label=f"Capa {index + 1}: {layer.soil.name}",
        )
        if cumulative < geometry.model_depth - 1e-9:
            contact = geometry.height - cumulative
            contact_mask = ground >= contact - 1e-9
            ax.plot(
                x[contact_mask],
                np.full(np.count_nonzero(contact_mask), contact),
                color="#7A6652",
                lw=1.1,
                ls=":",
                label="Contactos" if index == 0 else None,
            )
    ax.plot(x, ground, color="#17202A", lw=1.8, label="Terreno")
    water_mask = ground >= geometry.water_elevation - 1e-9
    ax.plot(
        x[water_mask],
        np.full(np.count_nonzero(water_mask), geometry.water_elevation),
        color="#2878B5",
        lw=1.6,
        ls="--",
        label=f"NF d={geometry.water_depth:.2f} m",
    )
    ax.axhline(bottom, color="#667085", lw=0.8, ls="--", label="Limite del modelo")
    xs = np.linspace(circle.entry_x, circle.exit_x, 350)
    ax.plot(xs, circle.base_elevation(xs), color="#D97706", lw=2.2, label=f"Bishop FS={critical.factor_of_safety:.3f}")
    for item in critical.evaluation.slices:
        ax.plot([item.x_mid, item.x_mid], [item.base_elevation, item.top_elevation], color="#65717B", lw=0.35, alpha=0.6)
    ax.scatter([circle.center_x], [circle.center_y], marker="+", color="#D97706", s=50)
    ax.set_xlabel("Distancia x (m)")
    ax.set_ylabel("Cota z (m)")
    ax.set_aspect("equal", adjustable="datalim")
    ax.grid(color="#E7E9ED", lw=0.6)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, 1.20), ncol=4, frameon=False, fontsize=7)
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _sensitivity_image(data: pd.DataFrame) -> BytesIO:
    fig, ax = plt.subplots(figsize=(8.1, 3.8), dpi=180)
    angles = data["Pendiente (deg)"].nunique()
    waters = data["Profundidad NF (m)"].nunique()
    if angles > 1 and waters > 1:
        pivot = data.pivot(index="Profundidad NF (m)", columns="Pendiente (deg)", values="FS Bishop")
        image = ax.imshow(pivot.values, origin="upper", aspect="auto", cmap="RdYlGn")
        ax.set_xticks(range(len(pivot.columns)), [f"{v:.1f}" for v in pivot.columns])
        ax.set_yticks(range(len(pivot.index)), [f"{v:.2f}" for v in pivot.index])
        ax.set_xlabel("Pendiente (deg)")
        ax.set_ylabel("Profundidad NF desde coronacion (m)")
        fig.colorbar(image, ax=ax, label="FS Bishop")
    else:
        variable = "Pendiente (deg)" if angles > 1 else "Profundidad NF (m)"
        ax.plot(data[variable], data["FS Fellenius"], "o-", label="Fellenius")
        ax.plot(data[variable], data["FS Bishop"], "o-", color="#D97706", label="Bishop simplificado")
        ax.set_xlabel(variable)
        ax.set_ylabel("Factor de seguridad")
        ax.legend(frameon=False)
    ax.grid(color="#E7E9ED", lw=0.6)
    fig.tight_layout()
    buffer = BytesIO()
    fig.savefig(buffer, format="png", bbox_inches="tight")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _set_style_font(style, name: str, size: float, color: str = "000000", bold: bool = False, italic: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = italic
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def _set_paragraph_border(style, color: str, size: int = 6, space: int = 1) -> None:
    p_pr = style._element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), color)
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    borders.append(bottom)
    p_pr.append(borders)


def _clear_paragraph_borders(style) -> None:
    p_pr = style._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def _add_style(document: Document, name: str, style_id: str):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style._element.set(qn("w:styleId"), style_id)
    return style


def _add_heading_numbering(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    for level in range(4):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), ".".join(f"%{index}" for index in range(1, level + 2)) + ".")
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        lvl.extend((start, num_fmt, lvl_text, suffix))
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _add_numbered_heading(document: Document, text: str, num_id: int, level: int = 1):
    paragraph = document.add_heading(text, level=level)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level - 1))
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, number))
    paragraph._p.get_or_add_pPr().insert(0, num_pr)
    return paragraph


def _configure_document(document: Document) -> int:
    section = document.sections[0]
    section.page_width = Cm(21.0009)
    section.page_height = Cm(29.7004)
    section.top_margin = Cm(2.5012)
    section.bottom_margin = Cm(2.5012)
    section.left_margin = Cm(3.0004)
    section.right_margin = Cm(3.0004)
    section.header_distance = Cm(1.2506)
    section.footer_distance = Cm(1.2506)
    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, "Open Sans", 9)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.07917
    geo_normal = _add_style(document, "GEO normal", "GEOnormal")
    _set_style_font(geo_normal, "Open Sans", 10)
    geo_normal.paragraph_format.space_before = Pt(6)
    geo_normal.paragraph_format.space_after = Pt(6)
    geo_normal.paragraph_format.line_spacing = 1.0
    geo_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, before, italic in (("Heading 1", 12, 18, False), ("Heading 2", 10, 12, True), ("Heading 3", 10, 6, False)):
        style = styles[name]
        _set_style_font(style, "Open Sans Condensed", size, GEO_BLUE if name != "Heading 3" else "000000", True, italic)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    styles["Heading 1"].font.all_caps = True
    _set_paragraph_border(styles["Heading 1"], GEO_BLUE)
    _set_style_font(styles["Title"], "Open Sans Condensed", 22, "000000", True)
    styles["Title"].paragraph_format.space_after = Pt(12)
    _clear_paragraph_borders(styles["Title"])
    figure_style = _add_style(document, "GEO pie de figura", "GEOpiedefigura")
    _set_style_font(figure_style, "Open Sans Light", 9, GEO_BLUE, True)
    figure_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_style.paragraph_format.space_before = Pt(6)
    figure_style.paragraph_format.space_after = Pt(12)
    table_title = _add_style(document, "GEO tabla titulo", "GEOtablatitulo")
    _set_style_font(table_title, "Open Sans Light", 9, GEO_BLUE, True)
    table_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_title.paragraph_format.space_after = Pt(6)
    table_title.paragraph_format.keep_with_next = True
    table_content = _add_style(document, "GEO tabla contenido", "GEOtablacontenido")
    _set_style_font(table_content, "Arial", 9)
    table_content.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_content.paragraph_format.line_spacing = 1.0
    list_style = styles["List Bullet"]
    _set_style_font(list_style, "Open Sans", 10)
    list_style.paragraph_format.space_before = Pt(3)
    list_style.paragraph_format.space_after = Pt(3)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Analisis preliminar de estabilidad de talud temporal | GeoTech")
    footer_run.font.name = "Open Sans Light"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GEO_BLUE)
    return _add_heading_numbering(document)


def build_report(
    project: ProjectData,
    geometry: SlopeGeometry,
    profile: SoilProfile,
    config: AnalysisConfig,
    result: SearchResult,
    sensitivity: pd.DataFrame | None = None,
) -> bytes:
    """Genera un DOCX en memoria listo para descarga."""
    document = Document()
    heading_num_id = _configure_document(document)
    title = document.add_paragraph(style="Title")
    title.add_run("Estabilidad de talud temporal")
    subtitle = document.add_paragraph(style="GEO normal")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.add_run(project.name or "Analisis sin titulo").bold = True
    subtitle.add_run(f"\n{project.location}" if project.location else "")
    subtitle.add_run(f"\nFecha: {datetime.now():%d/%m/%Y}")
    if project.author:
        subtitle.add_run(f"  |  Autor: {project.author}")
    document.add_paragraph(
        "Informe de calculo por equilibrio limite para superficies circulares. Los resultados son una ayuda al proyecto y no sustituyen la definicion del modelo geotecnico, las combinaciones de acciones ni las comprobaciones exigibles en el marco normativo aplicable.",
        style="GEO normal",
    )

    _add_numbered_heading(document, "Resumen de resultados", heading_num_id)
    governing = min(result.fellenius.factor_of_safety, result.bishop.factor_of_safety)
    status = "Cumple el objetivo introducido" if governing >= config.target_factor else "No alcanza el objetivo introducido"
    _add_table(
        document,
        ["Metodo", "FS minimo", "Centro x (m)", "Centro z (m)", "Radio (m)", "Evaluacion"],
        [
            [
                critical.method,
                f"{critical.factor_of_safety:.3f}",
                f"{critical.evaluation.circle.center_x:.2f}",
                f"{critical.evaluation.circle.center_y:.2f}",
                f"{critical.evaluation.circle.radius:.2f}",
                "FS >= objetivo" if critical.factor_of_safety >= config.target_factor else "FS < objetivo",
            ]
            for critical in (result.fellenius, result.bishop)
        ],
        "Resumen de las superficies criticas",
        1,
    )
    document.add_paragraph(
        f"Resultado gobernante: FS = {governing:.3f}. Factor objetivo definido por el usuario: {config.target_factor:.2f}. {status}. Se evaluaron {result.evaluated_surfaces} circulos, de los cuales {result.valid_surfaces} fueron geometricamente y numericamente validos.",
        style="GEO normal",
    )
    picture = document.add_picture(_section_image(geometry, profile, result), width=Cm(14.8))
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_caption(document, "Figura", 1, "Geometria y superficie critica obtenida con Bishop simplificado")

    document.add_page_break()
    _add_numbered_heading(document, "Datos de entrada", heading_num_id)
    _add_table(
        document,
        ["Parametro", "Valor", "Unidad / criterio"],
        [
            ["Altura H", f"{geometry.height:.2f}", "m"],
            ["Angulo del talud", f"{geometry.angle_deg:.2f}", "grados respecto a horizontal"],
            ["Desarrollo horizontal", f"{geometry.run:.2f}", "m"],
            ["Profundidad del nivel freatico", f"{geometry.water_depth:.2f}", "m desde coronacion"],
            ["Cota interna del nivel freatico", f"{geometry.water_elevation:.2f}", "m respecto al pie"],
            ["Profundidad del modelo", f"{geometry.model_depth:.2f}", "m desde coronacion"],
            ["Numero de dovelas", str(config.slices), "-"],
            ["Busqueda", config.search_level, "malla determinista"],
            ["Peso unitario del agua", f"{config.gamma_water:.2f}", "kN/m3"],
            ["Factor objetivo", f"{config.target_factor:.2f}", "definido por el usuario"],
        ],
        "Geometria y parametros de calculo",
        2,
    )
    soil_rows = []
    cumulative = 0.0
    for index, layer in enumerate(profile.layers, start=1):
        start_depth = cumulative
        cumulative += layer.thickness
        soil = layer.soil
        soil_rows.append(
            [
                str(index),
                soil.name,
                f"{layer.thickness:.2f}",
                f"{start_depth:.2f}-{cumulative:.2f}",
                f"{soil.gamma_dry:.2f}",
                f"{soil.gamma_sat:.2f}",
                f"{soil.cohesion:.2f}",
                f"{soil.friction_deg:.2f}",
            ]
        )
    _add_table(
        document,
        ["Capa", "Suelo", "Espesor", "Prof. desde coronacion", "gamma seco", "gamma sat.", "c'", "phi'"],
        soil_rows,
        "Parametros geotecnicos del perfil estratigrafico",
        3,
    )
    document.add_paragraph(
        f"El perfil contiene {len(profile.layers)} capas con un espesor acumulado de {profile.total_thickness:.2f} m. "
        f"El dominio de calculo alcanza {geometry.model_depth:.2f} m desde coronacion.",
        style="GEO normal",
    )

    _add_numbered_heading(document, "Metodo e hipotesis", heading_num_id)
    assumptions = [
        "Analisis bidimensional por metro de ancho, deformacion plana y mecanismo de rotura circular.",
        "Resistencia drenada Mohr-Coulomb con c' y phi' constantes dentro de cada estrato.",
        "Fellenius desprecia las fuerzas entre dovelas; Bishop simplificado satisface el equilibrio de momentos e itera el factor de seguridad.",
        "La profundidad del nivel freatico se mide desde coronacion. La presion intersticial es hidrostatica bajo el nivel horizontal y nula sobre este, sin modelar agua exterior ni fuerzas de filtracion adicionales.",
        "Los pesos se integran por estrato y condicion seca/saturada. La resistencia de la base de cada dovela corresponde al suelo atravesado por su punto medio.",
        "La superficie critica es el minimo dentro de la familia discreta explorada; una busqueda mas densa o una geometria real pueden producir un minimo inferior.",
    ]
    for item in assumptions:
        document.add_paragraph(item, style="List Bullet")

    has_sensitivity = sensitivity is not None and not sensitivity.empty
    if has_sensitivity:
        _add_numbered_heading(document, "Sensibilidad parametrica", heading_num_id)
        minimum_row = sensitivity.loc[sensitivity["FS Bishop"].idxmin()]
        maximum_row = sensitivity.loc[sensitivity["FS Bishop"].idxmax()]
        document.add_paragraph(
            f"Se analizaron {len(sensitivity)} combinaciones. Bishop varia entre {minimum_row['FS Bishop']:.3f} y {maximum_row['FS Bishop']:.3f}. El minimo corresponde a pendiente {minimum_row['Pendiente (deg)']:.1f} grados y profundidad del nivel freatico d={minimum_row['Profundidad NF (m)']:.2f} m desde coronacion.",
            style="GEO normal",
        )
        sensitivity_picture = document.add_picture(_sensitivity_image(sensitivity), width=Cm(14.8))
        sensitivity_picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_caption(document, "Figura", 2, "Sensibilidad del factor de seguridad frente a los parametros analizados")
        display = sensitivity.copy()
        if len(display) > 10:
            display = display.sort_values("FS Bishop").head(10)
            document.add_paragraph(
                "La tabla muestra las 10 combinaciones mas desfavorables; el CSV descargable conserva todos los casos.",
                style="GEO normal",
            )
        rows = [
            [f"{row['Pendiente (deg)']:.2f}", f"{row['Profundidad NF (m)']:.2f}", f"{row['FS Fellenius']:.3f}", f"{row['FS Bishop']:.3f}"]
            for _, row in display.iterrows()
        ]
        _add_table(
            document,
            ["Pendiente (deg)", "Prof. NF (m)", "FS Fellenius", "FS Bishop"],
            rows,
            "Resultados de sensibilidad parametrica",
            4,
        )

    _add_numbered_heading(document, "Limitaciones y recomendaciones", heading_num_id)
    document.add_paragraph(
        "Antes de adoptar el resultado deben contrastarse la estratigrafia, el regimen piezometrico, la representatividad de c' y phi', la posible respuesta no drenada, las sobrecargas, las excavaciones por fases y los mecanismos no circulares. El factor objetivo debe fijarse expresamente segun normativa, duracion de la situacion temporal, fiabilidad de la investigacion y consecuencias de fallo.",
        style="GEO normal",
    )
    if project.notes:
        _add_numbered_heading(document, "Observaciones del proyecto", heading_num_id)
        document.add_paragraph(project.notes, style="GEO normal")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
