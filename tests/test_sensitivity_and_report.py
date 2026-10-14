from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from slope_stability_app.calculations import run_sensitivity, search_critical_surfaces
from slope_stability_app.models import AnalysisConfig, ProjectData, SlopeGeometry, Soil, SoilLayer, SoilProfile, SweepConfig
from slope_stability_app.services import build_report


def _case():
    geometry = SlopeGeometry(6.0, 42.0, 4.0, 12.0)
    profile = SoilProfile((SoilLayer(Soil("Limo arenoso", 18.0, 20.0, 6.0, 29.0), 12.0),))
    config = AnalysisConfig(slices=18, search_level="Rapida", target_factor=1.25)
    return geometry, profile, config


def test_sensitivity_grid_has_expected_number_of_cases():
    geometry, profile, config = _case()
    sweep = SweepConfig(True, 35.0, 50.0, 3, True, 0.0, 5.0, 2)
    data = run_sensitivity(geometry, profile, config, sweep)
    assert len(data) == 6
    assert data["FS Bishop"].notna().all()
    assert set(data.columns) == {
        "Pendiente (deg)",
        "Profundidad NF (m)",
        "FS Fellenius",
        "FS Bishop",
        "Superficies validas",
    }


def test_report_is_valid_docx(tmp_path):
    geometry, profile, config = _case()
    result = search_critical_surfaces(geometry, profile, config)
    report = build_report(ProjectData("Caso de prueba", "Madrid", "Equipo geotecnico"), geometry, profile, config, result)
    path = tmp_path / "report.docx"
    path.write_bytes(report)
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Estabilidad de talud temporal" in text
    assert "Resumen de resultados" in text
    assert "Profundidad del nivel freatico" in table_text
    assert abs(document.sections[0].page_width - Cm(21.0009)) < Cm(0.01)
    assert abs(document.sections[0].page_height - Cm(29.7004)) < Cm(0.01)
    assert abs(document.sections[0].left_margin - Cm(3.0004)) < Cm(0.01)
    assert document.styles["GEO normal"].font.name == "Open Sans"
    assert document.styles["GEO pie de figura"].font.color.rgb == document.styles["Heading 1"].font.color.rgb
    headings = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
    assert headings
    assert all(paragraph._p.pPr.numPr is not None for paragraph in headings)
    instructions = " ".join(
        node.text or "" for node in document.element.body.iter(qn("w:instrText"))
    )
    assert "SEQ Tabla" in instructions
    assert "SEQ Figura" in instructions
    assert len(report) > 20_000
