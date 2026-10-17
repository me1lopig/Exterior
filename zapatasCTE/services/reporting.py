from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
import pandas as pd

from models.inputs import (
    CuModel,
    DrainageCondition,
    Foundation,
    InfluenceDepthMode,
    SettlementMode,
    SoilLayer,
    SoilProfile,
)


BLUE = "2164AF"
LIGHT_GRAY = "F2F2F2"
FONT = "Open Sans"
HEADING_FONT = "Open Sans Condensed"


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Pagina ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210.009)
    section.page_height = Mm(297.004)
    section.top_margin = Mm(25.012)
    section.right_margin = Mm(30.004)
    section.bottom_margin = Mm(25.012)
    section.left_margin = Mm(30.004)
    section.header_distance = Mm(12.506)
    section.footer_distance = Mm(12.506)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = doc.styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for level, size in ((1, 12), (2, 10), (3, 10)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.all_caps = level == 1
        style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else "000000")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(6)
        if level == 1:
            p_pr = style.element.get_or_add_pPr()
            borders = p_pr.find(qn("w:pBdr"))
            if borders is None:
                borders = OxmlElement("w:pBdr")
                p_pr.append(borders)
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), BLUE)
            borders.append(bottom)

    _add_page_number(section.footer.paragraphs[0])


def _add_key_value_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(6.0)
    table.columns[1].width = Cm(8.8)
    header = table.rows[0].cells
    header[0].text = "Parametro"
    header[1].text = "Valor"
    _set_repeat_table_header(table.rows[0])
    for cell in header:
        _shade_cell(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = FONT
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, (label, value) in enumerate(rows):
        row = table.add_row()
        _prevent_row_split(row)
        cells = row.cells
        cells[0].text = label
        cells[1].text = value
        if index % 2:
            for cell in cells:
                _shade_cell(cell, LIGHT_GRAY)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(9)
    _set_table_borders(table)


def _add_dataframe(doc: Document, frame: pd.DataFrame, formats: dict[str, str] | None = None) -> None:
    formats = formats or {}
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    for cell, column in zip(header, frame.columns):
        cell.text = str(column)
        _shade_cell(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.name = FONT
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    _set_repeat_table_header(table.rows[0])
    for row_index, (_, row) in enumerate(frame.iterrows()):
        table_row = table.add_row()
        _prevent_row_split(table_row)
        cells = table_row.cells
        for cell, column in zip(cells, frame.columns):
            value = row[column]
            if pd.isna(value):
                text = "-"
            elif column in formats:
                text = formats[column].format(value)
            elif isinstance(value, float):
                text = f"{value:.3f}"
            else:
                text = str(value)
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                _shade_cell(cell, LIGHT_GRAY)
            _set_cell_margins(cell, top=60, bottom=60)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(7.5)
    _set_table_borders(table)


def generate_docx_report(
    project: dict[str, str],
    soil: SoilProfile,
    foundation: Foundation,
    summary: pd.DataFrame,
    details: list[dict[str, Any]],
    serviceability: dict[str, Any] | None = None,
    layer_warnings: list[str] | None = None,
    bulb_coverage: dict[str, Any] | None = None,
) -> bytes:
    doc = Document()
    _configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Memoria geotecnica de cimentacion directa")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.add_run("Comprobaciones de hundimiento y servicio segun CTE DB SE C").bold = True
    doc.add_paragraph(
        "La memoria documenta los datos, hipotesis, resultados y limitaciones del "
        "analisis geotecnico. No incluye el dimensionamiento estructural de la cimentacion."
    )

    doc.add_heading("Identificacion", level=1)
    _add_key_value_table(
        doc,
        [
            ("Proyecto", project.get("name", "Sin denominacion")),
            ("Localizacion", project.get("location", "No indicada")),
            ("Autor o revision", project.get("author", "No indicado")),
            ("Observaciones", project.get("notes", "-")),
        ],
    )

    doc.add_heading("Modelo geotecnico", level=1)
    _add_key_value_table(
        doc,
        [
            ("Regimen", soil.drainage.value),
            ("Cohesion efectiva c'", f"{soil.c_effective_kpa:.2f} kPa"),
            ("Angulo de rozamiento phi'", f"{soil.phi_effective_deg:.2f} grados"),
            ("Resistencia no drenada cu0", f"{soil.cu0_kpa:.2f} kPa"),
            ("Gradiente de cu", f"{soil.cu_gradient_kpa_m:.2f} kPa/m"),
            ("Peso especifico aparente", f"{soil.gamma_apparent_kn_m3:.2f} kN/m3"),
            ("Peso especifico saturado", f"{soil.gamma_saturated_kn_m3:.2f} kN/m3"),
            ("Nivel freatico", f"{soil.water_table_depth_m:.2f} m"),
            ("Gradiente ascendente", f"{soil.upward_gradient:.3f}"),
        ],
    )

    doc.add_heading("Cimentacion", level=1)
    units = "por metro" if foundation.kind.value == "corrida" else "totales"
    _add_key_value_table(
        doc,
        [
            ("Tipologia", foundation.kind.value),
            ("Ancho B", f"{foundation.width_m:.3f} m"),
            ("Longitud L", f"{foundation.length_m:.3f} m"),
            ("Cota de apoyo D", f"{foundation.embedment_m:.3f} m"),
            ("Pendiente beta", f"{foundation.ground_slope_deg:.2f} grados"),
            ("Convencion de acciones", units),
        ],
    )

    if bulb_coverage:
        doc.add_heading("Cobertura de los bulbos de presiones", level=1)
        coverage_rows = []
        for check in bulb_coverage["checks"]:
            coverage_rows.append(
                {
                    "Control": check["label"],
                    "Desde m": check["required_top_m"],
                    "Hasta m": check["required_bottom_m"],
                    "Estado": "COMPLETA" if check["covered"] else "INSUFICIENTE",
                }
            )
        _add_dataframe(doc, pd.DataFrame(coverage_rows))
        if not bulb_coverage["all_covered"]:
            doc.add_paragraph(
                "No se acepta el asiento estratificado ni la geometria como solucion "
                "completa mientras existan intervalos del bulbo sin caracterizar."
            )
        doc.add_page_break()

    doc.add_heading("Resultados de las combinaciones", level=1)
    action_rows = []
    for detail in details:
        action = detail["actions"]
        action_rows.append(
            {
                "Caso": action.name,
                "V kN": action.vertical_kn,
                "HB kN": action.horizontal_b_kn,
                "HL kN": action.horizontal_l_kn,
                "MB kNm": action.moment_b_knm,
                "ML kNm": action.moment_l_knm,
            }
        )
    _add_dataframe(doc, pd.DataFrame(action_rows))
    doc.add_paragraph(
        "Las acciones anteriores incluyen el peso propio geotecnico configurado "
        "en la aplicacion. En zapata corrida se expresan por metro lineal."
    )

    summary_report = summary[
        [
            "combinacion",
            "situacion",
            "q_b_kpa",
            "q_d_kpa",
            "util_hundimiento",
            "cumple_geotecnia",
        ]
    ].rename(
        columns={
            "combinacion": "Caso",
            "situacion": "Situacion",
            "q_b_kpa": "qb kPa",
            "q_d_kpa": "qd kPa",
            "util_hundimiento": "Eta h",
            "cumple_geotecnia": "Estado ELU",
        }
    )
    summary_report["Estado ELU"] = summary_report["Estado ELU"].map(
        {True: "CUMPLE", False: "NO CUMPLE"}
    )
    _add_dataframe(doc, summary_report)

    doc.add_page_break()
    doc.add_heading("Desglose de hundimiento", level=1)
    for detail_index, detail in enumerate(details):
        if detail_index:
            doc.add_page_break()
        bearing = detail["bearing"]
        doc.add_heading(f"Combinacion {bearing.action_name}", level=2)
        rows = [
            ("Presion bruta qb", f"{bearing.pressure_gross_kpa:.2f} kPa"),
            ("Resistencia caracteristica qh", f"{bearing.characteristic_resistance_kpa:.2f} kPa"),
            ("Resistencia de calculo qd", f"{bearing.design_resistance_kpa:.2f} kPa"),
            ("Aprovechamiento", f"{bearing.utilization:.3f}"),
            ("Termino de cohesion", f"{bearing.terms_kpa['cohesion']:.2f} kPa"),
            ("Termino de sobrecarga", f"{bearing.terms_kpa['sobrecarga']:.2f} kPa"),
            ("Termino de peso especifico", f"{bearing.terms_kpa['peso_especifico']:.2f} kPa"),
            ("Coeficiente parcial gamma_R", f"{bearing.factors['gamma_R']:.2f}"),
        ]
        _add_key_value_table(doc, rows)
        factors = bearing.factors
        factor_table = pd.DataFrame(
            [
                {"Grupo": "Capacidad N", "c": factors["Nc"], "q": factors["Nq"], "gamma": factors["Ngamma"]},
                {"Grupo": "Forma s", "c": factors["sc"], "q": factors["sq"], "gamma": factors["sgamma"]},
                {"Grupo": "Profundidad d", "c": factors["dc"], "q": factors["dq"], "gamma": factors["dgamma"]},
                {"Grupo": "Inclinacion i", "c": factors["ic"], "q": factors["iq"], "gamma": factors["igamma"]},
                {"Grupo": "Talud t", "c": factors["tc"], "q": factors["tq"], "gamma": factors["tgamma"]},
            ]
        )
        _add_dataframe(doc, factor_table)
        for warning in bearing.warnings:
            doc.add_paragraph(f"Limitacion: {warning}")

    if serviceability:
        doc.add_page_break()
        doc.add_heading("Aptitud al servicio", level=1)
        service_rows = [
            (str(key).replace("_", " "), str(value))
            for key, value in serviceability.items()
        ]
        _add_key_value_table(doc, service_rows)

    doc.add_heading("Hipotesis y limitaciones", level=1)
    limitations = [
        "Los parametros introducidos deben ser valores caracteristicos representativos del estado limite.",
        "El modelo analitico de hundimiento supone un terreno equivalente homogeneo en la zona relevante.",
        "La comprobacion de contacto completo no sustituye una verificacion de equilibrio al vuelco con acciones separadas.",
        "No se comprueba la estabilidad global; se emiten avisos cuando el CTE exige un estudio especifico.",
        "No se incluye el dimensionamiento estructural de hormigon ni armaduras.",
    ]
    limitations.extend(layer_warnings or [])
    for item in limitations:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Referencias", level=1)
    doc.add_paragraph(
        "Codigo Tecnico de la Edificacion. Documento Basico DB SE C Cimientos, "
        "apartados 2.4, 4.2, 4.3, 4.4 y Anejo F."
    )

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_capacity_report(
    project: dict[str, str],
    soil: SoilProfile,
    bearing_layer_name: str,
    layers: list[SoilLayer],
    cu_profile: tuple[tuple[float, float], ...],
    foundation_template: Foundation,
    allowable_settlement_mm: float,
    influence_depth_m: float,
    capacity_grid: pd.DataFrame,
    action_summary: pd.DataFrame | None = None,
    action_details: list[dict[str, Any]] | None = None,
    settlement_mode: SettlementMode = SettlementMode.RIGID_CHARACTERISTIC,
    influence_depth_mode: InfluenceDepthMode = InfluenceDepthMode.MANUAL,
    include_subgrade_modulus: bool = False,
    grid_details: list[dict[str, Any]] | None = None,
) -> bytes:
    """Generate the B-L admissible-capacity report using the stored GEO style."""
    doc = Document()
    _configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Carga admisible de cimentaciones directas")
    subtitle = doc.add_paragraph()
    subtitle.add_run(
        "Hundimiento CTE DB SE-C y asiento elastico estratificado"
    ).bold = True
    doc.add_paragraph(
        "La carga admisible se obtiene para carga vertical centrada. Se distingue "
        "entre resistencia de calculo ELU y limitacion por asiento ELS. No se incluye "
        "el dimensionamiento estructural ni la comprobacion de deslizamiento."
    )

    doc.add_heading("Identificacion", level=1)
    _add_key_value_table(
        doc,
        [
            ("Proyecto", project.get("name", "Sin denominacion")),
            ("Localizacion", project.get("location", "No indicada")),
            ("Autor o revision", project.get("author", "No indicado")),
            ("Observaciones", project.get("notes", "-")),
        ],
    )

    doc.add_heading("Criterios de calculo", level=1)
    if soil.drainage is DrainageCondition.DRAINED:
        strength_description = "parametros efectivos c' y phi'"
    elif soil.cu_model is CuModel.LINEAR_CTE:
        strength_description = (
            f"cu=c0+mz; c0={soil.cu0_kpa:.2f} kPa; "
            f"m={soil.cu_gradient_kpa_m:.2f} kPa/m; limite 2c0"
        )
    elif soil.cu_model is CuModel.TABULATED:
        strength_description = "perfil tabulado cu(z), interpolacion sin extrapolacion"
    else:
        strength_description = f"cu constante = {soil.cu0_kpa:.2f} kPa"
    settlement_mode_text = {
        SettlementMode.RIGID_CHARACTERISTIC: "Steinbrenner multicapa, rigida en punto caracteristico",
        SettlementMode.FLEXIBLE_CENTER: "Steinbrenner multicapa, flexible en el centro",
        SettlementMode.FLEXIBLE_CORNER: "Steinbrenner multicapa, flexible en la esquina",
    }[settlement_mode]
    _add_key_value_table(
        doc,
        [
            ("Tipologia", foundation_template.kind.value),
            ("Cota de apoyo D", f"{foundation_template.embedment_m:.3f} m"),
            ("Regimen", soil.drainage.value),
            ("Resistencia para hundimiento", strength_description),
            ("Estrato representativo ELU", bearing_layer_name),
            ("Nivel freatico", f"{soil.water_table_depth_m:.3f} m"),
            ("Asiento admisible", f"{allowable_settlement_mm:.2f} mm"),
            ("Profundidad de influencia ZI", f"{influence_depth_m:.2f} m"),
            ("Criterio de ZI", influence_depth_mode.value),
            ("Metodo de asiento", settlement_mode_text),
            ("Termino resistente q0", "SI" if foundation_template.include_overburden_term else "NO"),
            ("Terreno sobre la base permanente", "SI" if foundation_template.overburden_permanent else "NO"),
            ("Posibles grietas de retraccion", "SI" if foundation_template.shrinkage_cracks else "NO"),
        ],
    )

    doc.add_heading("Modelo estratigrafico", level=1)
    layer_frame = pd.DataFrame(
        [
            {
                "Capa": layer.name,
                "Techo m": layer.top_depth_m,
                "Muro m": layer.bottom_depth_m,
                "c' kPa": layer.c_effective_kpa,
                "phi grados": layer.phi_effective_deg,
                "gamma": layer.gamma_apparent_kn_m3,
                "gamma sat": layer.gamma_saturated_kn_m3,
                "E_s MPa": layer.deformation_modulus_mpa,
                "nu": layer.poisson_ratio,
            }
            for layer in layers
        ]
    )
    _add_dataframe(doc, layer_frame)

    if cu_profile:
        doc.add_heading("Perfil de resistencia no drenada", level=2)
        _add_dataframe(
            doc,
            pd.DataFrame(cu_profile, columns=["Profundidad z m", "cu kPa"]),
        )

    doc.add_heading("Capacidad para la malla B-L", level=1)
    doc.add_paragraph(
        "qd ELU es la resistencia bruta de calculo frente a hundimiento. q ELS "
        "designa la presion que produce el asiento "
        "admisible; no es una resistencia de calculo. qadm es el menor de ambos "
        "valores. Nadm descuenta los pesos configurados en la aplicacion."
    )
    capacity_columns = [
        "B_m", "L_m", "q_d_kpa", "q_adm_asiento_neta_kpa",
        "q_adm_asiento_bruta_kpa", "q_limite_bruta_kpa",
        "Q_limite_estructura_kn", "criterio_gobernante",
        "perfil_cubre_bulbos",
    ]
    capacity_report = capacity_grid[capacity_columns].rename(
        columns={
            "B_m": "B m",
            "L_m": "L m",
            "q_d_kpa": "qd ELU kPa",
            "q_adm_asiento_neta_kpa": "q ELS neta kPa",
            "q_adm_asiento_bruta_kpa": "q ELS bruta kPa",
            "q_limite_bruta_kpa": "qadm kPa",
            "Q_limite_estructura_kn": "Nadm kN",
            "criterio_gobernante": "Criterio",
            "perfil_cubre_bulbos": "Perfil",
        }
    )
    capacity_report["Perfil"] = capacity_report["Perfil"].map(
        {True: "SI", False: "NO"}
    )
    capacity_report["Criterio"] = capacity_report["Criterio"].replace(
        {"Hundimiento ELU": "ELU", "Asiento ELS": "ELS"}
    )
    doc.add_heading("Presiones y axil admisible", level=2)
    _add_dataframe(doc, capacity_report)

    settlement_columns = ["B_m", "L_m", "ZI_m", "criterio_ZI"]
    if include_subgrade_modulus:
        settlement_columns.append("Ks_calculado_MN_m3")
    settlement_report = capacity_grid[settlement_columns].rename(columns={
        "B_m": "B m",
        "L_m": "L m",
        "ZI_m": "ZI m",
        "criterio_ZI": "Criterio ZI",
        "Ks_calculado_MN_m3": "Ks=q/s MN/m3",
    })
    doc.add_heading("Profundidad de influencia y balasto calculado", level=2)
    _add_dataframe(doc, settlement_report)

    if capacity_grid["KsB_CTE_MN_m3"].notna().any():
        doc.add_page_break()
        doc.add_heading("Balasto transformado desde Ksp30", level=2)
        doc.add_paragraph(
            "KsB se obtiene mediante E.6 para terreno cohesivo o E.7 para terreno "
            "granular. Para zapatas BxL se aplica despues E.8. En zapata corrida "
            "se adopta su limite L tendiendo a infinito, por lo que KsBL = KsB."
        )
        plate_report = capacity_grid[
            [
                "B_m", "L_m", "Ksp30_MN_m3", "tipo_terreno_balasto",
                "KsB_CTE_MN_m3", "KsBL_CTE_MN_m3",
                "ecuaciones_balasto_CTE",
            ]
        ].rename(columns={
            "B_m": "B m",
            "L_m": "L m",
            "Ksp30_MN_m3": "Ksp30 MN/m3",
            "tipo_terreno_balasto": "Terreno",
            "KsB_CTE_MN_m3": "KsB MN/m3",
            "KsBL_CTE_MN_m3": "KsBL MN/m3",
            "ecuaciones_balasto_CTE": "CTE",
        })
        _add_dataframe(doc, plate_report)

    incomplete = int((~capacity_grid["perfil_cubre_bulbos"]).sum())
    if incomplete:
        doc.add_paragraph(
            f"Aviso: {incomplete} geometria(s) no disponen de cobertura estratigrafica "
            "completa y no tienen carga admisible final."
        )

    if grid_details:
        valid = capacity_grid[capacity_grid["q_limite_bruta_kpa"].notna()]
        if not valid.empty:
            critical_index = int(valid["q_limite_bruta_kpa"].idxmin())
            detail = grid_details[critical_index]
            bearing = detail["bearing"]
            doc.add_heading("Auditoria de una geometria representativa", level=1)
            audit_intro = doc.add_paragraph(
                "Se documenta la geometria con menor presion admisible de la malla."
            )
            audit_intro.paragraph_format.keep_with_next = True
            _add_key_value_table(doc, [
                ("B*", f"{bearing.geometry.width_m:.3f} m"),
                ("L*", f"{bearing.geometry.length_m:.3f} m"),
                ("q0 disponible", f"{float(bearing.parameters['q0k_disponible_kpa']):.2f} kPa"),
                ("q0 resistente", f"{float(bearing.parameters['q0k_kpa']):.2f} kPa"),
                ("c o cu", f"{float(bearing.parameters['c_k_kpa']):.2f} kPa"),
                ("phi", f"{float(bearing.parameters['phi_k_deg']):.2f} grados"),
                ("qh", f"{bearing.characteristic_resistance_kpa:.2f} kPa"),
                ("qd ELU", f"{bearing.design_resistance_kpa:.2f} kPa"),
                ("Aplicacion gamma_R", str(bearing.parameters["aplicacion_gamma_R"])),
                ("Factores de profundidad", str(bearing.parameters["motivo_factores_profundidad"])),
            ])
            doc.add_heading("Factores de capacidad y correccion", level=2)
            factor_frame = pd.DataFrame([
                {"Grupo": "N", "c": bearing.factors["Nc"], "q": bearing.factors["Nq"], "gamma": bearing.factors["Ngamma"]},
                {"Grupo": "s", "c": bearing.factors["sc"], "q": bearing.factors["sq"], "gamma": bearing.factors["sgamma"]},
                {"Grupo": "d", "c": bearing.factors["dc"], "q": bearing.factors["dq"], "gamma": bearing.factors["dgamma"]},
                {"Grupo": "i", "c": bearing.factors["ic"], "q": bearing.factors["iq"], "gamma": bearing.factors["igamma"]},
                {"Grupo": "t", "c": bearing.factors["tc"], "q": bearing.factors["tq"], "gamma": bearing.factors["tgamma"]},
            ])
            _add_dataframe(doc, factor_frame)

    if action_summary is not None and not action_summary.empty:
        doc.add_page_break()
        doc.add_heading("Comprobacion opcional con acciones", level=1)
        action_report = action_summary[
            [
                "combinacion",
                "situacion",
                "B_eff_m",
                "L_eff_m",
                "q_b_kpa",
                "q_d_kpa",
                "util_hundimiento",
                "cumple_geotecnia",
            ]
        ].rename(
            columns={
                "combinacion": "Caso",
                "situacion": "Situacion",
                "B_eff_m": "B* m",
                "L_eff_m": "L* m",
                "q_b_kpa": "qb kPa",
                "q_d_kpa": "qd kPa",
                "util_hundimiento": "Eta h",
                "cumple_geotecnia": "Estado ELU",
            }
        )
        action_report["Estado ELU"] = action_report["Estado ELU"].map(
            {True: "CUMPLE", False: "NO CUMPLE"}
        )
        _add_dataframe(doc, action_report)
        doc.add_paragraph(
            "Esta comprobacion no incluye la resistencia frente al deslizamiento."
        )

    doc.add_heading("Formulaciones y alcance", level=1)
    for item in (
        "Hundimiento: CTE DB SE-C 4.3 y Anejo F.1.1, con geometria efectiva cuando existen excentricidades.",
        "Coeficiente parcial de hundimiento: en tensiones efectivas afecta a qh; en condiciones transitorias sin drenaje puede aplicarse solo al termino cohesivo conforme a F.1.1.2.5.",
        "Asiento: solucion elastica de Steinbrenner por superposicion de capas, con el termino de Poisson (1-nu-2nu^2); metodo analitico no normativo.",
        "Balasto: definicion ks=q/s y transformacion de Ksp30 mediante CTE DB SE-C Anejo E.5, expresiones E.6 a E.8.",
        "Presion limite: minimo entre qd ELU y la presion bruta limitada por el asiento admisible.",
        "El perfil debe cubrir D+1,5B para el control del bulbo de hundimiento y D+ZI para el asiento.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Limitaciones", level=1)
    for item in (
        "Los parametros del estrato portante deben ser valores caracteristicos representativos del mecanismo de rotura.",
        "Steinbrenner supone comportamiento elastico lineal y modulos compatibles con el nivel tensional de servicio.",
        "No se calculan consolidacion primaria, asiento secundario, estabilidad global ni dimensionamiento estructural.",
        "La equivalencia rectangular de una zapata circular para asientos conserva el area, pero no reproduce exactamente la respuesta axisimetrica.",
        "La transformacion de Ksp30 no se aplica a zapatas circulares y puede no representar estratos deformables situados fuera del bulbo del ensayo de placa.",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Referencia", level=1)
    doc.add_paragraph(
        "Codigo Tecnico de la Edificacion. Documento Basico DB SE-C Cimientos, "
        "apartados 2.4, 4.3, 4.4 y Anejos E y F."
    )

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
